"""End-to-end online exam code, scoring and export tenant isolation regression."""
import io
import json
import tempfile
import threading
import unittest
from pathlib import Path
from http.server import ThreadingHTTPServer
from urllib.request import Request,urlopen
from urllib.error import HTTPError
from openpyxl import load_workbook
from docx import Document
import app

class OnlineExamFlow(unittest.TestCase):
 def setUp(self):
  app.RATE.clear()
  self.temp=tempfile.TemporaryDirectory()
  self.old=(app.DB,app.TENANT_DIR,app.ADMIN_USER,app.ADMIN_PASSWORD)
  app.DB=Path(self.temp.name)/'app.db'
  app.TENANT_DIR=Path(self.temp.name)/'tenants';app.TENANT_DIR.mkdir()
  app.ADMIN_USER='admin_test'
  app.ADMIN_PASSWORD='TestAdminPass123'
  app.TENANT.set(None);app.init()
  self.http=ThreadingHTTPServer(('127.0.0.1',0),app.H)
  self.thread=threading.Thread(target=self.http.serve_forever,daemon=True);self.thread.start()
  status,credentials,headers=self.request('/api/login',{'username':'admin_test','password':'TestAdminPass123'})
  self.assertEqual(status,200)
  self.adminCookie=headers['Set-Cookie'].split(';',1)[0];self.csrf=credentials['csrf']
  app.TENANT.set(None)
  with app.db() as db:
   db.execute('INSERT INTO users(username,password_hash,role) VALUES(?,?,?)',('teacher_other',app.hashpw('TeacherOther123'),'teacher'))
  app.TENANT.set(app.tenant_for('teacher_other'));app.init();app.TENANT.set(None)
  status,d,h=self.request('/api/login',{'username':'teacher_other','password':'TeacherOther123','portal':'teacher'})
  self.assertEqual(status,200);self.otherCookie=h['Set-Cookie'].split(';',1)[0];self.otherCsrf=d['csrf']
 def tearDown(self):
  self.http.shutdown();self.http.server_close();self.thread.join(timeout=5)
  app.DB,app.TENANT_DIR,app.ADMIN_USER,app.ADMIN_PASSWORD=self.old
  app.TENANT.set(None);app.RATE.clear();self.temp.cleanup()
 def request(self,path,payload=None,cookie=None,csrf=None):
  headers={}
  if payload is not None:headers['Content-Type']='application/json'
  if cookie:headers['Cookie']=cookie
  if csrf:headers['X-CSRF-Token']=csrf
  req=Request(f'http://127.0.0.1:{self.http.server_port}'+path,data=json.dumps(payload).encode() if payload is not None else None,headers=headers)
  try:
   with urlopen(req,timeout=15) as r:
    raw=r.read();kind=r.headers.get('Content-Type','')
    return r.status,(json.loads(raw) if 'application/json' in kind else raw),r.headers
  except HTTPError as e:return e.code,json.loads(e.read()),e.headers
 def teacher(self,path,payload=None):
  return self.request(path,payload,self.adminCookie,self.csrf)
 def test_teacher_issued_exam_auto_grade_and_exports(self):
  questions=[{'content':'Hệ điều hành nào sau đây?','options':{'A':'Microsoft Word','B':'Windows','C':'PowerPoint','D':'Paint'},'answer':'B'},
   {'content':'Thiết bị nào xuất âm thanh?','options':{'A':'Chuột','B':'Loa','C':'Máy quét','D':'Bàn phím'},'answer':'B'}]
  status,published,_=self.teacher('/api/exams/word/publish',{'title':'Kiểm tra Tin học 7','subject':'Tin học','grade':'7','duration':20,'items':questions})
  self.assertEqual(status,201,published);eid=published['id']
  status,_,_=self.request('/api/online/share',{'exam_id':eid})
  self.assertEqual(status,401)
  status,_,_=self.request('/api/online/share',{'exam_id':eid},self.otherCookie,self.otherCsrf)
  self.assertEqual(status,404)
  status,share,_=self.teacher('/api/online/share',{'exam_id':eid})
  self.assertEqual(status,200,share);code=share['exam_code']
  status,_,_=self.request('/api/online/start',{'exam_code':'INVALID00000','student_name':'Học sinh thử','student_code':'HS001'})
  self.assertEqual(status,404)
  status,started,_=self.request('/api/online/start',{'exam_code':code,'student_name':'Học sinh thử','student_code':'HS001'})
  self.assertEqual(status,200,started)
  self.assertEqual(len(started['items']),2)
  self.assertNotIn('correct',str(started))
  app.TENANT.set(None)
  with app.db() as db:
   raw=db.execute('SELECT q.* FROM questions q JOIN exam_questions eq ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY eq.ord',(eid,)).fetchall()
  from online_exam import make_form
  hidden=make_form(raw,app.TENANT.get() or 0)
  # Inspect randomized issued form through server-stored seed, not the browser.
  with app.db() as db:run=db.execute('SELECT seed FROM online_exam_sessions WHERE token_hash=?',(__import__('online_exam').digest(started['attempt_token']),)).fetchone()
  forms=make_form(raw,int(run['seed']))
  answers={str(q['id']):q['correct'] for q in forms}
  status,done,_=self.request('/api/online/submit',{'attempt_token':started['attempt_token'],'answers':answers})
  self.assertEqual(status,200,done)
  self.assertEqual(done['score'],10.0)
  status,_,_=self.request('/api/online/submit',{'attempt_token':started['attempt_token'],'answers':answers})
  self.assertIn(status,(404,409))
  status,_,_=self.request('/api/online/start',{'exam_code':code,'student_name':'Học sinh thử','student_code':'HS001'})
  self.assertEqual(status,409)
  status,_,_=self.request('/api/online/results?exam_id='+str(eid),cookie=self.otherCookie)
  self.assertEqual(status,404)
  status,data,_=self.teacher('/api/online/results?exam_id='+str(eid))
  self.assertEqual(status,200,data);self.assertEqual(len(data['items']),1)
  self.assertEqual(data['items'][0]['student_code'],'HS001')
  status,raw,_=self.teacher('/api/online/results/export?exam_id='+str(eid)+'&format=xlsx')
  self.assertEqual(status,200)
  workbook=load_workbook(io.BytesIO(raw),read_only=True)
  self.assertEqual(workbook.active['F3'].value,10.0)
  status,raw,_=self.teacher('/api/online/results/export?exam_id='+str(eid)+'&format=docx')
  self.assertEqual(status,200)
  document=Document(io.BytesIO(raw))
  self.assertIn('Học sinh thử',' '.join(c.text for row in document.tables[0].rows for c in row.cells))
  status,replaced,_=self.teacher('/api/online/share',{'exam_id':eid})
  self.assertEqual(status,200,replaced)
  status,_,_=self.request('/api/online/start',{'exam_code':code,'student_name':'Học sinh hai','student_code':'HS002'})
  self.assertEqual(status,404)

if __name__=='__main__':unittest.main()
