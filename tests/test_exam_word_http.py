"""HTTP contract checks for teacher Word import and shuffled ZIP export."""
import base64
import io
import json
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from http.server import ThreadingHTTPServer
from urllib.request import Request,urlopen
from urllib.error import HTTPError
from docx import Document
import app

class WordExamHttpTest(unittest.TestCase):
 def setUp(self):
  app.RATE.clear()
  self.temp=tempfile.TemporaryDirectory()
  self.orig=(app.DB,app.TENANT_DIR,app.ADMIN_USER,app.ADMIN_PASSWORD)
  app.DB=Path(self.temp.name)/'app.db'
  app.TENANT_DIR=Path(self.temp.name)/'teachers';app.TENANT_DIR.mkdir()
  app.ADMIN_USER='admin_word';app.ADMIN_PASSWORD='StrongAdminTest123'
  app.TENANT.set(None);app.init()
  self.server=ThreadingHTTPServer(('127.0.0.1',0),app.H)
  self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
  status,body,headers=self.req('/api/login',{'username':'admin_word','password':'StrongAdminTest123'})
  self.assertEqual(status,200)
  self.csrf=body['csrf'];self.cookie=headers['Set-Cookie'].split(';',1)[0]
 def tearDown(self):
  self.server.shutdown();self.server.server_close();self.thread.join(timeout=5)
  app.DB,app.TENANT_DIR,app.ADMIN_USER,app.ADMIN_PASSWORD=self.orig
  app.TENANT.set(None);app.RATE.clear();self.temp.cleanup()
 def req(self,path,data=None,auth=False):
  headers={'Content-Type':'application/json'}
  if auth:headers.update({'Cookie':self.cookie,'X-CSRF-Token':self.csrf})
  request=Request('http://127.0.0.1:'+str(self.server.server_port)+path,
                  data=json.dumps(data).encode() if data is not None else None,headers=headers)
  try:
   with urlopen(request,timeout=20) as response:
    content=response.read()
    result=content if response.headers.get('Content-Type','').startswith('application/zip') else json.loads(content)
    return response.status,result,response.headers
  except HTTPError as error:return error.code,json.loads(error.read()),error.headers
 def test_admin_teacher_workflow_and_auth(self):
  doc=Document()
  for text in ('Câu 1. Hệ điều hành nào được sử dụng?','A. Windows','B. Word','C. Excel','D. PowerPoint','Đáp án: A',
               'Câu 2. Câu trả lời đúng là?','A. Sai','B. Đúng','C. Khác','D. Không','Đáp án: B'):
   doc.add_paragraph(text)
  output=io.BytesIO();doc.save(output)
  payload={'filename':'dethi.docx','file':base64.b64encode(output.getvalue()).decode('ascii')}
  status,_,_=self.req('/api/exams/word/preview',payload)
  self.assertEqual(status,401)
  status,data,_=self.req('/api/exams/word/preview',payload,True)
  self.assertEqual(status,200,data)
  self.assertEqual(data['count'],2)
  request={'title':'Kiểm tra Tin học','variants':8,'items':data['items']}
  status,_,_=self.req('/api/exams/word/export',request)
  self.assertEqual(status,401)
  status,result,_=self.req('/api/exams/word/export',request,True)
  self.assertEqual(status,200,result)
  with zipfile.ZipFile(io.BytesIO(result)) as archive:
   self.assertEqual(len(archive.namelist()),9)
   answers=Document(io.BytesIO(archive.read('Bang_dap_an.docx')))
   self.assertIn('Mã đề 008',' '.join(p.text for p in answers.paragraphs))
  request['items'][0]['answer']=''
  status,result,_=self.req('/api/exams/word/export',request,True)
  self.assertEqual(status,400,result)

if __name__=='__main__':unittest.main()
