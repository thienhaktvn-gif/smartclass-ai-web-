"""Test atomic attendance save and idempotent update with disposable SQLite."""
import json
import tempfile
import threading
import unittest
from pathlib import Path
from http.server import ThreadingHTTPServer
from urllib.request import Request,urlopen
from urllib.error import HTTPError
import app

class AttendanceBulkTest(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory()
  self.orig=(app.DB,app.TENANT_DIR,app.ADMIN_USER,app.ADMIN_PASSWORD)
  app.DB=Path(self.tmp.name)/'test.sqlite3'
  app.TENANT_DIR=Path(self.tmp.name)/'tenant'
  app.TENANT_DIR.mkdir()
  app.ADMIN_USER='att_teacher'
  app.ADMIN_PASSWORD='temporary-test-password'
  app.TENANT.set(None)
  app.init()
  self.http=ThreadingHTTPServer(('127.0.0.1',0),app.H)
  self.thread=threading.Thread(target=self.http.serve_forever,daemon=True)
  self.thread.start()
  code,data,cookie=self.call('/api/login',{'username':'att_teacher','password':'temporary-test-password'},False)
  self.assertEqual(code,200,data)
  self.cookie=cookie.split(';',1)[0];self.csrf=data['csrf']
  app.TENANT.set(None)
  with app.db() as conn:
   class_id=conn.execute('INSERT INTO classes(name,grade) VALUES(?,?)',('8A','8')).lastrowid
   conn.executemany('INSERT INTO students(code,name,class_id) VALUES(?,?,?)',[(f'S{i:04}',f'Học sinh {i}',class_id) for i in range(120)])
   self.ids=[r['id'] for r in conn.execute('SELECT id FROM students ORDER BY id')]
  app.TENANT.set(None)

 def tearDown(self):
  self.http.shutdown();self.http.server_close();self.thread.join(timeout=5)
  app.DB,app.TENANT_DIR,app.ADMIN_USER,app.ADMIN_PASSWORD=self.orig
  app.TENANT.set(None);self.tmp.cleanup()

 def call(self,path,data,authorized=True):
  headers={'Content-Type':'application/json'}
  if authorized:headers.update({'Cookie':self.cookie,'X-CSRF-Token':self.csrf})
  request=Request(f'http://127.0.0.1:{self.http.server_port}{path}',data=json.dumps(data).encode(),headers=headers)
  try:
   with urlopen(request,timeout=30) as response:return response.status,json.load(response),response.headers.get('Set-Cookie','')
  except HTTPError as error:return error.code,json.load(error),''

 def test_save_120_and_update_is_idempotent(self):
  items=[{'student_id':sid,'status':'present','note':''} for sid in self.ids]
  code,result,_=self.call('/api/attendance',{'day':'2026-10-24','items':items})
  self.assertEqual(code,200,result);self.assertEqual(result['saved'],120)
  items[0]['status']='absent';items[0]['note']='Có phép báo sau'
  code,result,_=self.call('/api/attendance',{'day':'2026-10-24','items':items})
  self.assertEqual(code,200,result);self.assertEqual(result['saved'],120)
  app.TENANT.set(None)
  with app.db() as conn:
   self.assertEqual(conn.execute('SELECT COUNT(*) FROM attendance').fetchone()[0],120)
   rec=conn.execute('SELECT status,note FROM attendance WHERE student_id=?',(self.ids[0],)).fetchone()
   self.assertEqual(rec['status'],'absent')
   self.assertEqual(rec['note'],'Có phép báo sau')

 def get(self,path,authorized=True):
  headers={'Cookie':self.cookie} if authorized else {}
  req=Request(f'http://127.0.0.1:{self.http.server_port}{path}',headers=headers)
  try:
   with urlopen(req,timeout=30) as response:return response.status,response.read(),response.headers.get('Content-Type','')
  except HTTPError as error:return error.code,error.read(),''

 def test_month_report_and_xlsx_export(self):
  import io
  from openpyxl import load_workbook
  app.TENANT.set(None)
  with app.db() as conn:class_id=conn.execute('SELECT id FROM classes WHERE name=?',('8A',)).fetchone()['id']
  endpoint=f'/api/attendance/report?class_id={class_id}&month=2026-10'
  status,raw,_=self.get(endpoint,False)
  self.assertEqual(status,401)
  status,raw,_=self.get(endpoint)
  self.assertEqual(status,200)
  before=json.loads(raw)
  self.assertEqual(before['saved_days'],0)
  self.assertFalse(before['days'][23]['saved'])
  self.assertEqual(before['students'][0]['recorded'],0)
  records=[{'student_id':i,'status':'present'} for i in self.ids]
  records[0]['status']='absent'
  records[1]['status']='late'
  records[2]['status']='excused'
  status,data,_=self.call('/api/attendance',{'day':'2026-10-24','items':records})
  self.assertEqual(status,200,data)
  status,raw,_=self.get(endpoint)
  report=json.loads(raw);self.assertEqual(report['saved_days'],1)
  saved=report['days'][23]
  self.assertTrue(saved['saved'])
  self.assertEqual((saved['present'],saved['absent'],saved['late'],saved['excused'],saved['unrecorded']),(117,1,1,1,0))
  self.assertFalse(report['days'][22]['saved'])
  self.assertEqual(sum(r['recorded'] for r in report['students']),120)
  records[0]['status']='present'
  self.call('/api/attendance',{'day':'2026-10-24','items':records})
  status,raw,_=self.get(endpoint);report=json.loads(raw)
  self.assertEqual(report['saved_days'],1)
  self.assertEqual(report['days'][23]['present'],118)
  status,raw,mime=self.get(f'/api/attendance/report/export?class_id={class_id}&month=2026-10')
  self.assertEqual(status,200);self.assertIn('spreadsheet',mime)
  book=load_workbook(io.BytesIO(raw),read_only=True)
  self.assertEqual(book.sheetnames,['Tong hop theo ngay','Chuyen can hoc sinh','Chi tiet diem danh'])
  self.assertEqual(book.worksheets[0]['B26'].value,'Đã lưu')
  self.assertEqual(book.worksheets[0]['B25'].value,'Chưa điểm danh')
  self.assertEqual(book.worksheets[2].max_row,121)
  for path in (f'/api/attendance/report?class_id={class_id}&month=2026-13','/api/attendance/report?class_id=999999&month=2026-02'):
   status,_,_=self.get(path)
   self.assertEqual(status,400)

 def test_reject_duplicate_students_without_partial_write(self):
  rows=[{'student_id':self.ids[0],'status':'present'},{'student_id':self.ids[0],'status':'absent'}]
  code,_,_=self.call('/api/attendance',{'day':'2026-10-24','items':rows})
  self.assertEqual(code,400)
  app.TENANT.set(None)
  with app.db() as conn:self.assertEqual(conn.execute('SELECT COUNT(*) FROM attendance').fetchone()[0],0)

if __name__=='__main__':unittest.main()
