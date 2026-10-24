"""Exercise Word importer, shuffled answer labels and generated DOCX files."""
import io
import re
import unittest
import zipfile
from docx import Document
from docx.shared import Inches
from exam_word import parse_word,make_variants,export_package

def word_bytes(lines):
 doc=Document()
 for line in lines:doc.add_paragraph(line)
 out=io.BytesIO();doc.save(out);return out.getvalue()

class WordExamTest(unittest.TestCase):
 def source(self):
  return word_bytes(['ĐỀ KIỂM TRA',
   'Câu 1. Python có câu lệnh nào để kiểm tra điều kiện?',
   'A. for','B. if','C. print','D. input',
   'Câu 2. Ký hiệu phép cộng trong Python là gì?',
   'A. +','B. -','C. /','D. *',
   'Câu 3. Số nào là số nguyên?',
   'A. 3.5','B. "3"','C. 3','D. True',
   'Câu 4. Hàm nào xuất dữ liệu?',
   'A. input','B. str','C. int','D. print',
   'ĐÁP ÁN','1.B 2.A 3.C 4.D'])
 def test_import_and_variants_answer_mapping(self):
  questions=parse_word(self.source())
  self.assertEqual([q['answer'] for q in questions],list('BACD'))
  for count in (2,4,8):
   variants=make_variants(questions,count)
   self.assertEqual(len(variants),count)
   self.assertEqual(len({tuple((q['source_number'],tuple(q['options'].values())) for q in variant) for variant in variants}),count)
   originals={q['number']:q for q in questions}
   for variant in variants:
    self.assertEqual(len(variant),len(questions))
    for q in variant:
     src=originals[q['source_number']]
     self.assertEqual(q['options'][q['answer']],src['options'][src['answer']])
   raw=export_package('Kiểm tra thử môn Tin học',questions,count)
   with zipfile.ZipFile(io.BytesIO(raw)) as archive:
    self.assertEqual(len(archive.namelist()),count+1)
    self.assertIn('Bang_dap_an.docx',archive.namelist())
    answers=Document(io.BytesIO(archive.read('Bang_dap_an.docx')))
    answer_text=' '.join(p.text for p in answers.paragraphs)
    for i in range(1,count+1):
     code=f'{i:03d}'
     doc=Document(io.BytesIO(archive.read(f'Ma_de_{code}.docx')))
     text=' '.join(p.text for p in doc.paragraphs)
     self.assertIn('MÃ ĐỀ: '+code,text)
     expected=' '.join(f'{j}.{q["answer"]}' for j,q in enumerate(variants[0],1))
     self.assertIn('Mã đề '+code,answer_text)
 def test_incomplete_answer_and_duplicates_rejected(self):
  bad=word_bytes(['Câu 1. Mặt trời mọc từ hướng nào?','A. Đông','B. Tây','C. Nam','D. Bắc'])
  with self.assertRaisesRegex(ValueError,'Thiếu đáp án'):parse_word(bad)
  conflict=word_bytes(['Câu 1. Nội dung nào đúng?','A. Một','B. Hai','C. Ba','D. Bốn','Đáp án: A','ĐÁP ÁN','1.B'])
  with self.assertRaisesRegex(ValueError,'mâu thuẫn'):parse_word(conflict)
 def test_reject_drawing_and_tables(self):
  import base64
  from tempfile import TemporaryDirectory
  with TemporaryDirectory() as path:
   file=path+'/tiny.png'
   with open(file,'wb') as out:out.write(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a7mYAAAAASUVORK5CYII='))
   d=Document();d.add_paragraph('Câu 1. Hình nào đúng?');d.add_picture(file,width=Inches(1))
   b=io.BytesIO();d.save(b)
   with self.assertRaisesRegex(ValueError,'hình ảnh'):parse_word(b.getvalue())
  doc=Document();doc.add_table(rows=2,cols=2);b=io.BytesIO();doc.save(b)
  with self.assertRaisesRegex(ValueError,'bảng Word'):parse_word(b.getvalue())

if __name__=='__main__':unittest.main()
