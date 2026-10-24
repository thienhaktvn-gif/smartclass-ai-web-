"""Secure, teacher-scoped online multiple-choice assessment helpers."""
import io
import hashlib
import random
import re
import secrets
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from docx import Document
from docx.shared import Pt

LETTERS='ABCD'
def digest(value):return hashlib.sha256(value.encode('utf-8')).hexdigest()
def create_exam_code():
 alphabet='ABCDEFGHJKLMNPQRSTUVWXYZ23456789'
 return ''.join(secrets.choice(alphabet) for _ in range(14))
def valid_student(name,code):
 return 2<=len(name)<=80 and bool(re.fullmatch(r'[A-Z0-9_-]{3,32}',code))
def make_form(questions,seed):
 rng=random.Random(seed);rows=list(questions);rng.shuffle(rows)
 result=[]
 for q in rows:
  order=list(LETTERS);rng.shuffle(order)
  options=[{'key':LETTERS[i],'text':q[order[i].lower()]} for i in range(4)]
  result.append({'id':q['id'],'content':q['content'],'options':options,'correct':'ABCD'[order.index(q['answer'])]})
 return result
def score_form(items,answers):
 if not isinstance(answers,dict):raise ValueError('Câu trả lời không hợp lệ.')
 ids={str(q['id']) for q in items}
 if len(answers)>len(items) or any(k not in ids or v not in LETTERS for k,v in answers.items()):
  raise ValueError('Câu trả lời không đúng mẫu đề đã cấp.')
 correct=sum(answers.get(str(q['id']))==q['correct'] for q in items)
 return {'correct':correct,'score':round(10*correct/len(items),2) if items else 0}
def online_result_rows(c,exam_id):
 rows=c.execute('SELECT r.student_code,a.student_name,a.score,a.correct,a.total,a.feedback,a.created FROM online_exam_results r JOIN attempts a ON a.id=r.attempt_id WHERE r.exam_id=? ORDER BY a.created,r.student_code',(exam_id,)).fetchall()
 return [dict(row) for row in rows]
def safe_excel(value):
 value=str(value or '')
 return "'"+value if value.startswith(('=','+','-','@')) else value
def make_results_export(exam,rows,kind):
 title=str(exam['title'])
 if kind=='xlsx':
  wb=Workbook();sheet=wb.active;sheet.title='Bang diem'
  sheet.append(['BẢNG ĐIỂM KIỂM TRA',safe_excel(title)])
  sheet.append(['STT','Mã học sinh','Họ và tên','Số câu đúng','Tổng số câu','Điểm /10','Nhận xét','Ngày nộp'])
  for n,r in enumerate(rows,1):
   sheet.append([n,safe_excel(r['student_code']),safe_excel(r['student_name']),r['correct'],r['total'],r['score'],safe_excel(r['feedback']),safe_excel(r['created'])])
  for cell in sheet[2]:
   cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='274690')
  sheet.freeze_panes='D3'
  sheet.auto_filter.ref=f'A2:H{max(2,sheet.max_row)}'
  for col,width in {'A':10,'B':20,'C':32,'D':16,'E':16,'F':14,'G':45,'H':25}.items():sheet.column_dimensions[col].width=width
  output=io.BytesIO();wb.save(output);return output.getvalue()
 if kind=='docx':
  doc=Document();style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(11)
  doc.add_heading('BẢNG ĐIỂM KIỂM TRA',0);doc.add_paragraph('Đề: '+title)
  table=doc.add_table(rows=1,cols=6);table.style='Table Grid'
  for cell,heading in zip(table.rows[0].cells,['STT','Mã học sinh','Họ tên','Đúng/Tổng','Điểm /10','Nhận xét']):cell.text=heading
  for n,r in enumerate(rows,1):
   cells=table.add_row().cells
   for cell,value in zip(cells,[n,r['student_code'],r['student_name'],str(r['correct'])+'/'+str(r['total']),r['score'],r['feedback']]):cell.text=str(value)
  doc.add_paragraph('Số lượt nộp hợp lệ: '+str(len(rows)))
  output=io.BytesIO();doc.save(output);return output.getvalue()
 raise ValueError('Chỉ hỗ trợ Excel hoặc Word.')
