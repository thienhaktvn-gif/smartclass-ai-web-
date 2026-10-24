import os, json, sqlite3, secrets, hashlib, hmac, random, re, urllib.request, urllib.error, base64, uuid, io
from pathlib import Path
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse
from http.cookies import SimpleCookie
from datetime import datetime, timezone
from contextvars import ContextVar
import time, threading, calendar
from openpyxl import Workbook, load_workbook

SUBJECT_CATALOG=['Toán', 'Ngữ văn', 'Tiếng Anh', 'Vật lí', 'Hóa học', 'Sinh học', 'Lịch sử', 'Địa lí', 'Giáo dục công dân', 'Tin học', 'Công nghệ', 'Khoa học tự nhiên', 'Âm nhạc', 'Mĩ thuật', 'Giáo dục thể chất', 'Hoạt động trải nghiệm']

ROOT=Path(__file__).parent
DB=Path(os.environ.get('SMARTCLASS_DB',str(ROOT/'smartclass.db')))
DB.parent.mkdir(parents=True,exist_ok=True)
SECRET=os.environ.get('SMARTCLASS_SECRET','')
ADMIN_PASSWORD=os.environ.get('SMARTCLASS_ADMIN_PASSWORD','')
ADMIN_USER=os.environ.get('SMARTCLASS_ADMIN_USER','giaovien')
AI_KEY=os.environ.get('GEMINI_API_KEY','')
AI_MODEL=os.environ.get('GEMINI_MODEL','gemini-3.5-flash-lite')
IMAGE_MODEL=os.environ.get('GEMINI_IMAGE_MODEL','gemini-nano-banana-2.1')
VIDEO_MODEL=os.environ.get('GEMINI_VIDEO_MODEL','gemini-omni-1.1-flash')
MEDIA=ROOT/'media';MEDIA.mkdir(exist_ok=True)
PUBLIC_URL=os.environ.get('PUBLIC_URL','')
SESSION_LIFETIME=43200
TENANT=ContextVar('smartclass_tenant', default=None)
TENANT_DIR=DB.parent/'smartclass_teachers'
TENANT_DIR.mkdir(parents=True,exist_ok=True)
RATE={}
RATE_LOCK=threading.Lock()
def throttle(key,limit,window):
 now=time.monotonic()
 with RATE_LOCK:
  a=[t for t in RATE.get(key,[]) if now-t<window]
  if len(a)>=limit:return False
  a.append(now);RATE[key]=a
  return True

def tenant_for(name):
 if not name or name==ADMIN_USER:return None
 digest=hashlib.sha256(name.encode('utf8')).hexdigest()
 if os.environ.get('SMARTCLASS_STORAGE_BACKEND')=='postgres':return 'teacher_'+digest
 return TENANT_DIR/(digest+'.sqlite3')


def db():
 if os.environ.get('SMARTCLASS_STORAGE_BACKEND')=='postgres':
  from pgadapter import Connection
  url=os.environ.get('DATABASE_URL','')
  if not url:raise RuntimeError('DATABASE_URL missing for PostgreSQL backend')
  return Connection(url, TENANT.get() or 'public')
 c=sqlite3.connect(TENANT.get() or DB,timeout=15); c.row_factory=sqlite3.Row;c.execute('PRAGMA foreign_keys=ON');return c

def init():
 with db() as c:
  c.executescript('''CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,username TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,role TEXT NOT NULL,created TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE TABLE IF NOT EXISTS login_sessions(token_hash TEXT PRIMARY KEY,username TEXT NOT NULL,csrf TEXT NOT NULL,expires INTEGER NOT NULL);
  CREATE TABLE IF NOT EXISTS online_exam_links(code_hash TEXT PRIMARY KEY,teacher_username TEXT NOT NULL,exam_id INTEGER NOT NULL,active INTEGER NOT NULL DEFAULT 1,created TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE TABLE IF NOT EXISTS online_exam_sessions(token_hash TEXT PRIMARY KEY,code_hash TEXT NOT NULL,teacher_username TEXT NOT NULL,exam_id INTEGER NOT NULL,student_name TEXT NOT NULL,student_code TEXT NOT NULL,seed INTEGER NOT NULL,expires INTEGER NOT NULL);
  CREATE TABLE IF NOT EXISTS password_reset_required(username TEXT PRIMARY KEY, issued INTEGER NOT NULL);
  CREATE TABLE IF NOT EXISTS classes(id INTEGER PRIMARY KEY,name TEXT NOT NULL UNIQUE,grade TEXT DEFAULT '');
  CREATE TABLE IF NOT EXISTS students(id INTEGER PRIMARY KEY,code TEXT UNIQUE NOT NULL,name TEXT NOT NULL,class_id INTEGER REFERENCES classes(id));
  CREATE TABLE IF NOT EXISTS questions(id INTEGER PRIMARY KEY,subject TEXT NOT NULL,grade TEXT NOT NULL,topic TEXT NOT NULL,content TEXT NOT NULL,a TEXT NOT NULL,b TEXT NOT NULL,c TEXT NOT NULL,d TEXT NOT NULL,answer TEXT NOT NULL CHECK(answer IN ('A','B','C','D')),difficulty TEXT NOT NULL DEFAULT 'Nhận biết',explanation TEXT NOT NULL DEFAULT '',origin TEXT DEFAULT 'manual');
  CREATE TABLE IF NOT EXISTS exams(id INTEGER PRIMARY KEY,title TEXT NOT NULL,subject TEXT NOT NULL,grade TEXT NOT NULL,duration INTEGER NOT NULL DEFAULT 45,published INTEGER NOT NULL DEFAULT 0,created TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE TABLE IF NOT EXISTS exam_questions(exam_id INTEGER REFERENCES exams(id) ON DELETE CASCADE,question_id INTEGER REFERENCES questions(id),ord INTEGER NOT NULL,PRIMARY KEY(exam_id,question_id));
  CREATE TABLE IF NOT EXISTS attempts(id INTEGER PRIMARY KEY,exam_id INTEGER NOT NULL REFERENCES exams(id),student_name TEXT NOT NULL,access_code_hash TEXT NOT NULL,form_seed INTEGER NOT NULL,score REAL NOT NULL,correct INTEGER NOT NULL,total INTEGER NOT NULL,answers TEXT NOT NULL,feedback TEXT NOT NULL,created TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE INDEX IF NOT EXISTS attempts_exam_idx ON attempts(exam_id);
  CREATE TABLE IF NOT EXISTS online_exam_results(attempt_id INTEGER PRIMARY KEY,session_hash TEXT UNIQUE NOT NULL,exam_id INTEGER NOT NULL,student_code TEXT NOT NULL,submitted TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(exam_id,student_code));
   CREATE TABLE IF NOT EXISTS attendance(id INTEGER PRIMARY KEY,student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,day TEXT NOT NULL,status TEXT NOT NULL CHECK(status IN ('present','absent','late','excused')),note TEXT DEFAULT '',UNIQUE(student_id,day));
   CREATE TABLE IF NOT EXISTS attendance_days(class_id INTEGER NOT NULL,day TEXT NOT NULL,PRIMARY KEY(class_id,day));
  CREATE TABLE IF NOT EXISTS demo_meta(k TEXT PRIMARY KEY,v TEXT NOT NULL);
  CREATE TABLE IF NOT EXISTS games(id INTEGER PRIMARY KEY,title TEXT NOT NULL,topic TEXT NOT NULL,kind TEXT NOT NULL,question_ids TEXT NOT NULL,created TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE TABLE IF NOT EXISTS lesson_media(id INTEGER PRIMARY KEY,title TEXT NOT NULL,kind TEXT NOT NULL,content TEXT NOT NULL,filename TEXT DEFAULT '',created TEXT DEFAULT CURRENT_TIMESTAMP);''')
  row=c.execute('SELECT id,role FROM users WHERE username=?',(ADMIN_USER,)).fetchone()
  if TENANT.get() is None and row and row['role']!='admin':
   c.execute('UPDATE users SET role=? WHERE id=?',('admin',row['id']))
  if TENANT.get() is None and not row and ADMIN_PASSWORD:
   salt=secrets.token_hex(16);v=hashlib.pbkdf2_hmac('sha256',ADMIN_PASSWORD.encode(),bytes.fromhex(salt),240000).hex()
   c.execute('INSERT INTO users(username,password_hash,role) VALUES(?,?,?)',(ADMIN_USER,salt+':'+v,'admin'))
  if TENANT.get() is None and c.execute('SELECT COUNT(*) FROM questions').fetchone()[0]==0 and os.getenv('SMARTCLASS_DEMO','0')=='1':
   rows=[('Tin học','9','Cấu trúc rẽ nhánh','Trong Python, từ khóa nào bắt đầu câu lệnh điều kiện?','for','if','def','while','B','Nhận biết','if được dùng để kiểm tra điều kiện.'),('Tin học','9','Vòng lặp','Hàm nào tạo dãy số nguyên trong Python?','range','print','input','str','A','Nhận biết','range tạo các giá trị cho vòng lặp.'),('Tin học','9','Kiểu dữ liệu','Kiểu dữ liệu nào biểu diễn Đúng/Sai?','str','list','bool','int','C','Nhận biết','bool có hai giá trị True và False.'),('Tin học','9','Thuật toán','Tìm kiếm nhị phân yêu cầu điều gì?','Mảng đảo ngược','Dữ liệu được sắp xếp','Tất cả phần tử trùng nhau','Chỉ số chẵn','B','Thông hiểu','Tìm kiếm nhị phân cần dãy đã sắp xếp.'),('Tin học','9','Mảng','Chỉ số đầu tiên của list Python là bao nhiêu?','1','-1','0','2','C','Nhận biết','List được đánh chỉ số từ 0.'),('Tin học','9','Độ phức tạp','Thuật toán duyệt lần lượt n phần tử có thời gian O(?)','1','log n','n','n²','C','Thông hiểu','Duyệt n phần tử tuyến tính là O(n).')]
   for r in rows:c.execute('INSERT INTO questions(subject,grade,topic,content,a,b,c,d,answer,difficulty,explanation,origin) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(*r,'demo'))

def gemini(parts,model=None,timeout=70,config=None):
 if not AI_KEY:raise RuntimeError('Chưa cấu hình GEMINI_API_KEY trên máy chủ.')
 model=model or AI_MODEL
 if not re.fullmatch(r'[a-zA-Z0-9_.-]{3,80}',model):raise RuntimeError('Tên model không hợp lệ.')
 url='https://generativelanguage.googleapis.com/v1beta/models/'+model+':generateContent'
 payload={'contents':[{'role':'user','parts':parts}]}
 if config:payload['generationConfig']=config
 req=urllib.request.Request(url,data=json.dumps(payload).encode(),headers={'x-goog-api-key':AI_KEY,'Content-Type':'application/json'},method='POST')
 try:
  with urllib.request.urlopen(req,timeout=timeout) as r:return json.load(r)
 except urllib.error.HTTPError as e:
  if e.code in (401,403):raise RuntimeError('Gemini từ chối khóa API hoặc quyền model.')
  if e.code==429:raise RuntimeError('Gemini đã hết hạn mức hoặc bị giới hạn tốc độ.')
  if e.code==503:raise RuntimeError('Gemini đang quá tải (HTTP 503). Hãy thử lại sau ít phút.')
  raise RuntimeError(f'Gemini trả lỗi HTTP {e.code}. Kiểm tra model và hạn mức.')
 except (urllib.error.URLError,TimeoutError):raise RuntimeError('Không kết nối được Gemini.')

def gemini_interaction(model, prompt, media_type, timeout=120):
 if not AI_KEY:raise RuntimeError('Chưa cấu hình GEMINI_API_KEY trên Render.')
 if not re.fullmatch(r'[a-zA-Z0-9_.-]{3,80}', model):raise RuntimeError('Tên model không hợp lệ.')
 if media_type not in ('image','video'):raise ValueError('Loại nội dung không hợp lệ')
 payload={'model':model,'input':prompt,'response_format':{'type':media_type}}
 if media_type=='video':payload['response_format']['resolution']='360p'
 req=urllib.request.Request('https://generativelanguage.googleapis.com/v1beta/interactions', data=json.dumps(payload).encode(),headers={'x-goog-api-key':AI_KEY,'Content-Type':'application/json'},method='POST')
 try:
  with urllib.request.urlopen(req,timeout=timeout) as response: result=json.load(response)
 except urllib.error.HTTPError as exc:
  try: detail=json.loads(exc.read(8192).decode()).get('error',{}).get('message','')
  except Exception:detail=''
  if exc.code==429:raise RuntimeError('Google API giới hạn tạo '+('ảnh' if media_type=='image' else 'video')+' (HTTP 429). Kiểm tra quota model trong Google AI Studio. '+str(detail)[:220])
  if exc.code in (401,403):raise RuntimeError('API Key không được cấp quyền tạo '+('ảnh' if media_type=='image' else 'video')+' (HTTP '+str(exc.code)+').')
  if exc.code==404:raise RuntimeError('Model '+model+' không khả dụng cho dự án hoặc endpoint (HTTP 404).')
  raise RuntimeError('Gemini '+media_type+' HTTP '+str(exc.code)+': '+str(detail)[:250])
 except (urllib.error.URLError,TimeoutError) as exc:raise RuntimeError('Hết thời gian hoặc không kết nối được dịch vụ tạo '+media_type+'.') from exc
 if result.get('status') not in (None,'completed'):
  raise RuntimeError('Google đang xử lý tác vụ '+media_type+'; phản hồi hiện chưa có tệp hoàn chỉnh. Trạng thái: '+str(result.get('status'))[:50])
 for step in result.get('steps',[]):
  for item in step.get('content',[]):
   if item.get('type')==media_type and item.get('data'):
    mime=item.get('mime_type','image/png' if media_type=='image' else 'video/mp4')
    try:raw=base64.b64decode(item['data'],validate=True)
    except Exception:raise RuntimeError('Google trả nội dung base64 không hợp lệ.')
    return raw,mime
 raise RuntimeError('Model không trả '+media_type+' trực tiếp. Kiểm tra quyền truy cập và cấu trúc phản hồi API.')

def gemini_text(prompt):
 r=gemini([{'text':prompt}],config={'responseMimeType':'application/json'})
 parts=[p.get('text','') for c in r.get('candidates',[]) for p in c.get('content',{}).get('parts',[])]
 if not parts:raise RuntimeError('Gemini không trả nội dung văn bản.')
 return json.loads(''.join(parts).strip().removeprefix('```json').removeprefix('```').removesuffix('```').strip())

def safe_game(row):
 return {'id':row['id'],'title':row['title'],'topic':row['topic'],'kind':row['kind'],'count':len(json.loads(row['question_ids'])),'created':row['created']}

def hashpw(p):
 salt=secrets.token_hex(16);return salt+':'+hashlib.pbkdf2_hmac('sha256',p.encode(),bytes.fromhex(salt),240000).hex()
def checkpw(p,v):
 try:
  salt,he=v.split(':');return hmac.compare_digest(hashlib.pbkdf2_hmac('sha256',p.encode(),bytes.fromhex(salt),240000).hex(),he)
 except:return False

def public_question(q):
 return {k:q[k] for k in ('id','content','a','b','c','d','difficulty')}
def shuffled(qs,seed):
 rng=random.Random(seed);a=list(qs);rng.shuffle(a);out=[]
 for q in a:
  opts=[{'label':z,'text':q[z.lower()]} for z in 'ABCD'];rng.shuffle(opts)
  out.append({'id':q['id'],'content':q['content'],'options':[{'key':'ABCD'[i],'text':o['text']} for i,o in enumerate(opts)],'_correct':'ABCD'[next(i for i,o in enumerate(opts) if o['label']==q['answer'])]})
 return out

def attendance_month_report(c,class_id,month):
 """Count persisted statuses only: missing attendance is not present."""
 if not re.fullmatch(r'[0-9]{4}-(0[1-9]|1[0-2])',month):raise ValueError('Tháng không hợp lệ (YYYY-MM).')
 year,mon=map(int,month.split('-'))
 if year<2000 or year>2100:raise ValueError('Năm báo cáo không hợp lệ.')
 cl=c.execute('SELECT id,name,grade FROM classes WHERE id=?',(class_id,)).fetchone()
 if not cl:raise LookupError('Không tìm thấy lớp học.')
 days=calendar.monthrange(year,mon)[1];start=f'{month}-01';end=f'{month}-{days:02}'
 students=[dict(r) for r in c.execute('SELECT id,code,name FROM students WHERE class_id=? ORDER BY name,code',(class_id,))]
 records=[dict(r) for r in c.execute('SELECT a.student_id,a.day,a.status,a.note FROM attendance a JOIN students s ON s.id=a.student_id WHERE s.class_id=? AND a.day BETWEEN ? AND ? ORDER BY a.day,s.name',(class_id,start,end))]
 saved={r['day'] for r in c.execute('SELECT day FROM attendance_days WHERE class_id=? AND day BETWEEN ? AND ?',(class_id,start,end))}
 saved.update(r['day'] for r in records)
 daily={f'{month}-{n:02}':{'day':f'{month}-{n:02}','saved':f'{month}-{n:02}' in saved,'present':0,'absent':0,'late':0,'excused':0,'unrecorded':len(students)} for n in range(1,days+1)}
 counts={r['id']:{'student_id':r['id'],'code':r['code'],'name':r['name'],'present':0,'absent':0,'late':0,'excused':0,'recorded':0} for r in students}
 for r in records:
  day=daily.get(r['day']);student=counts.get(r['student_id'])
  if not day or not student or r['status'] not in ('present','absent','late','excused'):continue
  day[r['status']]+=1;day['unrecorded']-=1
  student[r['status']]+=1;student['recorded']+=1
 return {'class':dict(cl),'month':month,'days':list(daily.values()),'students':list(counts.values()),'records':records,'saved_days':len(saved)}

class H(BaseHTTPRequestHandler):
 def log_message(self,fmt,*args):print('HTTP:',fmt%args)
 def sendj(self,code,data,headers=None):
  b=json.dumps(data,ensure_ascii=False).encode();self.send_response(code);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(b)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.send_header('X-Frame-Options','DENY');self.send_header('Referrer-Policy','same-origin');
  if headers:
   for k,v in headers.items():self.send_header(k,v)
  self.end_headers();self.wfile.write(b)
 def err(self,code,msg):self.sendj(code,{'error':msg})
 def body(self):
  try:
   n=int(self.headers.get('Content-Length','0'))
   if n>2500000 or n<1:raise ValueError()
   v=json.loads(self.rfile.read(n));return v if isinstance(v,dict) else None
  except:return None
 def session(self):
  try:
   cookie=SimpleCookie();cookie.load(self.headers.get('Cookie',''))
   raw=cookie['scsession'].value
   if not re.fullmatch(r'[A-Za-z0-9_-]{30,128}',raw):return None
   token_hash=hashlib.sha256(raw.encode('ascii')).hexdigest()
   context=TENANT.set(None)
   try:
    with db() as c:
     row=c.execute('SELECT s.username,s.csrf,s.expires,CASE WHEN r.username IS NULL THEN 0 ELSE 1 END must_change FROM login_sessions s LEFT JOIN password_reset_required r ON r.username=s.username WHERE s.token_hash=?',(token_hash,)).fetchone()
   finally:TENANT.reset(context)
   if row and int(row['expires'])>int(time.time()):
    return {'username':row['username'],'csrf':row['csrf'],'must_change':bool(row['must_change'])}
  except (KeyError,ValueError,sqlite3.Error):
   return None
  return None
 def auth(self):
  u=self.session()
  if not u or self.headers.get('X-CSRF-Token','')!=u['csrf']:self.err(401,'Cần đăng nhập hợp lệ.');return None
  return u
 def do_GET(self):
  ss=self.session(); TENANT.set(tenant_for(ss['username']) if ss else None)
  p=urlparse(self.path).path
  if ss and ss['must_change'] and p.startswith('/api/') and p not in ('/api/me','/api/health','/api/registration-status','/api/subjects'):
   return self.err(403,'Bạn cần đổi mật khẩu tạm trước khi sử dụng hệ thống.')
  if p=='/api/registration-status':return self.sendj(200,{'enabled':os.environ.get('SMARTCLASS_ALLOW_REGISTRATION','0')=='1' and (os.environ.get('SMARTCLASS_STORAGE_BACKEND')=='postgres' and bool(os.environ.get('DATABASE_URL')) or os.environ.get('SMARTCLASS_PERSISTENT','0')=='1' and DB.is_absolute() and str(DB).startswith('/var/data/'))})
  if p=='/api/subjects':return self.sendj(200,{'items':SUBJECT_CATALOG,'label':'Danh mục môn học dùng chung'})
  if p=='/api/health':return self.sendj(200,{'ok':True,'aiConfigured':bool(AI_KEY),'aiProvider':'Gemini','imageConfigured':bool(AI_KEY),'videoConfigured':bool(AI_KEY),'storage':('postgres' if os.getenv('SMARTCLASS_STORAGE_BACKEND')=='postgres' else 'sqlite'),'persistent':(os.getenv('SMARTCLASS_STORAGE_BACKEND')=='postgres' and bool(os.getenv('DATABASE_URL')) or os.getenv('SMARTCLASS_PERSISTENT','0')=='1')})
  if p in ('/','/index.html') or p.startswith('/auth/') or p in ('/dashboard','/teacher','/student','/planner','/games','/attendanceReport','/exams'):
   b=(ROOT/'index.html').read_bytes();self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.send_header('Content-Length',str(len(b)));self.send_header('X-Content-Type-Options','nosniff');self.end_headers();return self.wfile.write(b)
  if p in ('/app.js','/style.css','/ui-motion.js','/learning-addons.js','/portal-enhance.js','/multisubject.js','/accounts.js','/admin.js','/attendance-report.js','/word-exam.js','/online-exam-ui.js','/onboarding-fix.js','/showcase-demo.js','/gemini-studio.js','/chatbox.js','/bank16.js','/question-pages.js','/planner-ai.js','/school-admin.js'):
   b=(ROOT/p[1:]).read_bytes();self.send_response(200);self.send_header('Content-Type','text/javascript; charset=utf-8' if p.endswith('.js') else 'text/css; charset=utf-8');self.send_header('Content-Length',str(len(b)));self.end_headers();return self.wfile.write(b)
  if p=='/api/me':
   s=self.session()
   return self.sendj(200,{'user':s['username'],'csrf':s['csrf'],'role':'admin' if s['username']==ADMIN_USER else 'teacher','mustChangePassword':s['must_change']} if s else {'user':None,'role':'guest'})
  if p=='/api/admin/teachers':
   s=self.session()
   if not s:return self.err(401,'Cần đăng nhập.')
   if s['username']!=ADMIN_USER:return self.err(403,'Chỉ quản trị viên được xem danh sách tài khoản.')
   context=TENANT.set(None)
   try:
    with db() as c:
     rows=c.execute('SELECT username,role,created FROM users ORDER BY id DESC').fetchall()
     return self.sendj(200,{'items':[{'username':r['username'],'role':'admin' if r['username']==ADMIN_USER else 'teacher','created':r['created']} for r in rows],'total':len(rows)})
   finally:TENANT.reset(context)
  if p=='/api/exams/public':
   with db() as c:rr=c.execute('SELECT id,title,subject,grade,duration FROM exams WHERE published=1 ORDER BY id DESC').fetchall()
   return self.sendj(200,{'items':[dict(x) for x in rr]})
  if p.startswith('/api/exams/public/'):
   try:eid=int(p.rsplit('/',1)[-1])
   except:return self.err(400,'Mã đề không hợp lệ.')
   with db() as c:
    e=c.execute('SELECT id,title,subject,grade,duration FROM exams WHERE id=? AND published=1',(eid,)).fetchone()
    if not e:return self.err(404,'Đề chưa được công bố.')
    qs=c.execute('SELECT q.* FROM exam_questions eq JOIN questions q ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY eq.ord',(eid,)).fetchall()
   return self.sendj(200,{'exam':dict(e),'count':len(qs)})
  if p=='/api/games/public':
   with db() as c:rows=c.execute('SELECT * FROM games ORDER BY id DESC LIMIT 100').fetchall()
   return self.sendj(200,{'items':[safe_game(r) for r in rows]})
  if p.startswith('/api/games/public/'):
   try:gid=int(p.rsplit('/',1)[-1])
   except:return self.err(400,'Mã trò chơi sai.')
   with db() as c:
    g=c.execute('SELECT * FROM games WHERE id=?',(gid,)).fetchone()
    if not g:return self.err(404,'Không tìm thấy trò chơi.')
    ids=json.loads(g['question_ids'])
    q=c.execute('SELECT * FROM questions WHERE id IN ('+','.join('?' for _ in ids)+')',ids).fetchall() if ids else []
   rng=random.SystemRandom();q=list(q);rng.shuffle(q)
   return self.sendj(200,{'game':safe_game(g),'items':[public_question(x) for x in q]})
  if p.startswith('/media/'):
   name=p.split('/')[-1]
   if not re.fullmatch(r'[a-f0-9]{32}\.(png|jpg|webp|mp4)',name):return self.err(404,'Không tìm thấy hình ảnh.')
   s=self.session()
   if not s:return self.err(401,'Hãy đăng nhập giáo viên.')
   target=MEDIA/name
   if not target.is_file():return self.err(404,'Không tìm thấy hình ảnh.')
   b=target.read_bytes();self.send_response(200);self.send_header('Content-Type',{'png':'image/png','jpg':'image/jpeg','webp':'image/webp','mp4':'video/mp4'}[name.rsplit('.',1)[1]]);self.send_header('Content-Length',str(len(b)));self.send_header('Cache-Control','private, no-store');self.end_headers();return self.wfile.write(b)
  if p=='/api/online/results/export':
   if not ss:return self.err(401,'Cần đăng nhập giáo viên.')
   from online_exam import make_results_export,online_result_rows
   from urllib.parse import parse_qs
   q=parse_qs(urlparse(self.path).query)
   try:exam_id=int(q.get('exam_id',['0'])[0])
   except:return self.err(400,'Mã đề không hợp lệ.')
   kind=q.get('format',['xlsx'])[0]
   if kind not in ('xlsx','docx'):return self.err(400,'Chỉ hỗ trợ Excel hoặc Word.')
   with db() as c:
    exam=c.execute('SELECT id,title,subject,grade FROM exams WHERE id=?',(exam_id,)).fetchone()
    if not exam:return self.err(404,'Không tìm thấy đề.')
    rows=online_result_rows(c,exam_id)
   data=make_results_export(dict(exam),rows,kind)
   mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if kind=='xlsx' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
   self.send_response(200);self.send_header('Content-Type',mime)
   self.send_header('Content-Disposition',f'attachment; filename="Bang_diem_de_{exam_id}.{kind}"')
   self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','private, no-store')
   self.send_header('X-Content-Type-Options','nosniff');self.end_headers();return self.wfile.write(data)
  if p=='/api/attendance/report/export':
   if not self.session():return self.err(401,'Vui lòng đăng nhập giáo viên.')
   from urllib.parse import parse_qs
   q=parse_qs(urlparse(self.path).query)
   try:
    cid=int(q.get('class_id',[''])[0]);month=q.get('month',[''])[0]
    with db() as c:report=attendance_month_report(c,cid,month)
   except (ValueError,LookupError) as exc:return self.err(400,str(exc))
   wb=Workbook();daily=wb.active;daily.title='Tong hop theo ngay'
   daily.append(['BÁO CÁO CHUYÊN CẦN',report['class']['name'],'Tháng '+month])
   daily.append(['Ngày','Tình trạng','Có mặt','Vắng','Đi muộn','Có phép','Chưa ghi nhận'])
   for item in report['days']:
    daily.append([item['day'],'Đã lưu' if item['saved'] else 'Chưa điểm danh',item['present'] if item['saved'] else None,item['absent'] if item['saved'] else None,item['late'] if item['saved'] else None,item['excused'] if item['saved'] else None,item['unrecorded'] if item['saved'] else None])
   individual=wb.create_sheet('Chuyen can hoc sinh')
   individual.append(['Mã học sinh','Họ tên','Có mặt','Vắng','Đi muộn','Có phép','Số ngày ghi nhận'])
   # Force student-supplied text to literal Excel cells, never formulas.
   def safe(value):
    value=str(value or '')
    return "'"+value if value.startswith(('=','+','-','@')) else value
   for student in report['students']:
    individual.append([safe(student['code']),safe(student['name']),student['present'],student['absent'],student['late'],student['excused'],student['recorded']])
   detail=wb.create_sheet('Chi tiet diem danh')
   detail.append(['Ngày','Mã học sinh','Họ tên','Trạng thái','Ghi chú'])
   names={student['student_id']:student for student in report['students']}
   labels={'present':'Có mặt','absent':'Vắng','late':'Đi muộn','excused':'Có phép'}
   for record in report['records']:
    person=names.get(record['student_id'],{})
    detail.append([record['day'],safe(person.get('code','')),safe(person.get('name','')),labels.get(record['status'],record['status']),safe(record['note'])])
   from openpyxl.styles import Font,PatternFill
   for sheet,header in ((daily,2),(individual,1),(detail,1)):
    sheet.freeze_panes=f'A{header+1}'
    sheet.auto_filter.ref=f'A{header}:{sheet.cell(row=max(header,sheet.max_row),column=sheet.max_column).coordinate}'
    for cell in sheet[header]:
     cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='274690')
    for column in sheet.columns:
     cells=list(column)
     sheet.column_dimensions[cells[0].column_letter].width=min(44,max(16,max(len(str(cell.value or '')) for cell in cells[:100])+3))
   output=io.BytesIO();wb.save(output);data=output.getvalue()
   self.send_response(200);self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
   self.send_header('Content-Disposition',f'attachment; filename="Chuyen_can_lop_{cid}_{month}.xlsx"')
   self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','private, no-store');self.send_header('X-Content-Type-Options','nosniff');self.end_headers();return self.wfile.write(data)
  if p=='/api/students/export':
   if not self.session():return self.err(401,'Vui lòng đăng nhập giáo viên.')
   query=urlparse(self.path).query
   from urllib.parse import parse_qs
   try:cid=int(parse_qs(query).get('class_id',['0'])[0])
   except:return self.err(400,'Mã lớp không hợp lệ.')
   with db() as c:
    cl=c.execute('SELECT name,grade FROM classes WHERE id=?',(cid,)).fetchone()
    if not cl:return self.err(404,'Không tìm thấy lớp.')
    rows=c.execute('SELECT code,name FROM students WHERE class_id=? ORDER BY code',(cid,)).fetchall()
   wb=Workbook();ws=wb.active;ws.title='Danh sach hoc sinh';ws.append(['Mã học sinh','Họ và tên','Lớp'])
   for r in rows:ws.append([r['code'],r['name'],cl['name']])
   ws.column_dimensions['A'].width=22;ws.column_dimensions['B'].width=35;ws.column_dimensions['C'].width=25
   f=io.BytesIO();wb.save(f);data=f.getvalue();self.send_response(200)
   self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
   self.send_header('Content-Disposition',f'attachment; filename="smartclass_lop_{cid}.xlsx"')
   self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','no-store');self.end_headers();return self.wfile.write(data)
  if p.startswith('/api/'):
   s=self.session()
   if not s:return self.err(401,'Vui lòng đăng nhập giáo viên.')
   with db() as c:
    if p=='/api/classes':
     rows=c.execute('SELECT c.*,COUNT(s.id) student_count FROM classes c LEFT JOIN students s ON s.class_id=c.id GROUP BY c.id ORDER BY c.id DESC').fetchall()
     return self.sendj(200,{'items':[dict(x) for x in rows]})
    if p=='/api/students':
     return self.sendj(200,{'items':[dict(x) for x in c.execute('SELECT s.*,c.name class_name FROM students s LEFT JOIN classes c ON c.id=s.class_id ORDER BY c.name,s.name')]})
    if p=='/api/attendance/report':
     from urllib.parse import parse_qs
     q=parse_qs(urlparse(self.path).query)
     try:report=attendance_month_report(c,int(q.get('class_id',[''])[0]),q.get('month',[''])[0])
     except (ValueError,LookupError) as exc:return self.err(400,str(exc))
     report.pop('records')
     return self.sendj(200,report)
    if p=='/api/attendance':
     from urllib.parse import parse_qs
     q=parse_qs(urlparse(self.path).query)
     day=q.get('day',[''])[0];cid=q.get('class_id',['0'])[0]
     if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',day):return self.err(400,'Ngày không hợp lệ.')
     try:cid=int(cid)
     except:return self.err(400,'Lớp không hợp lệ.')
     items=c.execute('SELECT s.id student_id,s.code,s.name,a.status,a.note FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND a.day=? WHERE s.class_id=? ORDER BY s.name',(day,cid)).fetchall()
     return self.sendj(200,{'items':[dict(x) for x in items], 'day':day})
    if p=='/api/games':return self.sendj(200,{'items':[safe_game(r) for r in c.execute('SELECT * FROM games ORDER BY id DESC LIMIT 100')]})
    if p=='/api/media':return self.sendj(200,{'items':[dict(r) for r in c.execute('SELECT * FROM lesson_media ORDER BY id DESC LIMIT 100')]})
    if p=='/api/overview':
     data={k:c.execute('SELECT COUNT(*) FROM '+table).fetchone()[0] for k,table in [('questions','questions'),('exams','exams'),('attempts','attempts'),('students','students')]}
     return self.sendj(200,data)
    if p=='/api/questions/page':
     from urllib.parse import parse_qs
     args=parse_qs(urlparse(self.path).query)
     try:
      page=max(1,min(100000,int(args.get('page',['1'])[0])))
      size=max(1,min(100,int(args.get('size',['30'])[0])))
     except (ValueError,TypeError):return self.err(400,'Trang và số lượng phải là số nguyên.')
     filters=[];values=[]
     for key in ('subject','grade'):
      v=str(args.get(key,[''])[0]).strip()[:100]
      if v:filters.append(key+'=?');values.append(v)
     term=str(args.get('q',[''])[0]).strip()[:120]
     if term:
      filters.append('(content LIKE ? OR topic LIKE ?)')
      values.extend(['%'+term+'%','%'+term+'%'])
     where=' WHERE '+' AND '.join(filters) if filters else ''
     count=c.execute('SELECT COUNT(*) FROM questions'+where,values).fetchone()[0]
     items=[dict(row) for row in c.execute('SELECT * FROM questions'+where+' ORDER BY id DESC LIMIT ? OFFSET ?',values+[size,(page-1)*size])]
     return self.sendj(200,{'items':items,'page':page,'size':size,'total':count,'pages':(count+size-1)//size})
    if p=='/api/questions':return self.sendj(200,{'items':[dict(x) for x in c.execute('SELECT * FROM questions ORDER BY id DESC LIMIT 500')]})
    if p=='/api/exams':return self.sendj(200,{'items':[dict(x) for x in c.execute('SELECT e.*, COUNT(eq.question_id) num_questions FROM exams e LEFT JOIN exam_questions eq ON e.id=eq.exam_id GROUP BY e.id ORDER BY e.id DESC')]})
    if p=='/api/online/results':
     from online_exam import online_result_rows
     from urllib.parse import parse_qs
     q=parse_qs(urlparse(self.path).query)
     try:exam_id=int(q.get('exam_id',['0'])[0])
     except:return self.err(400,'Mã đề không hợp lệ.')
     exam=c.execute('SELECT id,title FROM exams WHERE id=?',(exam_id,)).fetchone()
     if not exam:return self.err(404,'Không tìm thấy đề kiểm tra.')
     return self.sendj(200,{'exam':dict(exam),'items':online_result_rows(c,exam_id)})
    if p=='/api/attempts':return self.sendj(200,{'items':[dict(x) for x in c.execute('SELECT a.id,a.exam_id,e.title exam,a.student_name,a.score,a.correct,a.total,a.feedback,a.created FROM attempts a JOIN exams e ON e.id=a.exam_id ORDER BY a.id DESC LIMIT 300')]})
   return self.err(404,'Không tồn tại API.')
  return self.err(404,'Không tìm thấy trang.')
 def do_POST(self):
  ss=self.session(); TENANT.set(tenant_for(ss['username']) if ss else None)
  p=urlparse(self.path).path;d=self.body()
  if ss and ss['must_change'] and p not in ('/api/password/change','/api/logout'):
   return self.err(403,'Bạn cần đổi mật khẩu tạm trước khi sử dụng hệ thống.')
  if d is None:return self.err(400,'Dữ liệu JSON không hợp lệ hoặc quá lớn.')
  if p=='/api/register':
   if not (os.environ.get('SMARTCLASS_ALLOW_REGISTRATION','0')=='1' and (os.environ.get('SMARTCLASS_STORAGE_BACKEND')=='postgres' and bool(os.environ.get('DATABASE_URL')) or os.environ.get('SMARTCLASS_PERSISTENT','0')=='1' and DB.is_absolute() and str(DB).startswith('/var/data/'))):return self.err(403,'Đăng ký chỉ được bật khi lưu trữ bền vững đã được cấu hình và xác minh.')
   address=self.client_address[0]
   name=str(d.get('username','')).strip().lower(); pw=str(d.get('password',''))
   if not re.fullmatch(r'[a-z][a-z0-9_.-]{4,31}',name):return self.err(400,'Tên đăng nhập gồm 5–32 ký tự không dấu, bắt đầu bằng chữ cái.')
   if name==ADMIN_USER:return self.err(409,'Tên đăng nhập không khả dụng.')
   if len(pw)<12 or len(pw)>128 or not re.search('[A-Z]',pw) or not re.search('[a-z]',pw) or not re.search('[0-9]',pw):return self.err(400,'Mật khẩu cần 12–128 ký tự, gồm chữ hoa, chữ thường và chữ số.')
   # Invalid form submissions should not consume the signup limit.
   if not throttle('register:'+address,5,3600):return self.err(429,'Đã gửi quá nhiều đăng ký hợp lệ trong một giờ. Vui lòng thử lại sau.')
   TENANT.set(None)
   with db() as c:
    if c.execute('SELECT 1 FROM users WHERE username=?',(name,)).fetchone():return self.err(409,'Tên đăng nhập đã được sử dụng.')
    c.execute('INSERT INTO users(username,password_hash,role) VALUES(?,?,?)',(name,hashpw(pw),'teacher'))
   TENANT.set(None)
   with db() as c:
    stored=c.execute('SELECT password_hash FROM users WHERE username=?',(name,)).fetchone()
   if not stored or not checkpw(pw,stored['password_hash']):
    return self.err(503,'Không thể xác nhận tài khoản đã được lưu. Vui lòng thử lại.')
   try:
    TENANT.set(tenant_for(name));init()
   finally:TENANT.set(None)
   return self.sendj(201,{'ok':True,'message':'Tạo tài khoản thành công. Mời đăng nhập.'})
  if p=='/api/login':
   if not throttle('login:'+self.client_address[0],15,900):return self.err(429,'Có quá nhiều lần thử đăng nhập. Vui lòng thử lại sau.')
   name=str(d.get('username','')).strip().lower()[:100];pw=str(d.get('password',''))
   TENANT.set(None)
   with db() as c:u=c.execute('SELECT * FROM users WHERE username=?',(name,)).fetchone()
   if not u or not checkpw(pw,u['password_hash']):return self.err(401,'Sai tài khoản hoặc mật khẩu.')
   actual_role='admin' if name==ADMIN_USER else 'teacher'
   portal=str(d.get('portal',actual_role)).strip().lower()
   if portal not in ('teacher','admin'):return self.err(400,'Cổng đăng nhập không hợp lệ.')
   if portal!=actual_role:return self.err(403,'Tài khoản không có quyền đăng nhập tại cổng đã chọn.')
   tok=secrets.token_urlsafe(32);csrf=secrets.token_urlsafe(24)
   context=TENANT.set(None)
   try:
    with db() as c:
     c.execute('DELETE FROM login_sessions WHERE expires<=?',(int(time.time()),))
     c.execute('INSERT INTO login_sessions(token_hash,username,csrf,expires) VALUES(?,?,?,?)',(hashlib.sha256(tok.encode('ascii')).hexdigest(),name,csrf,int(time.time())+SESSION_LIFETIME))
   finally:TENANT.reset(context)
   secure='; Secure' if self.headers.get('X-Forwarded-Proto','')=='https' or PUBLIC_URL.startswith('https://') else ''
   return self.sendj(200,{'user':name,'csrf':csrf},{'Set-Cookie':f'scsession={tok}; HttpOnly; SameSite=Lax; Path=/; Max-Age=43200{secure}'})
  if p=='/api/admin/password-reset':
   user=self.auth()
   if not user:return
   if user['username']!=ADMIN_USER:return self.err(403,'Chỉ quản trị viên được đặt lại mật khẩu.')
   target=str(d.get('username','')).strip().lower()
   if not re.fullmatch(r'[a-z][a-z0-9_.-]{4,31}',target) or target==ADMIN_USER:
    return self.err(400,'Chọn một tài khoản giáo viên hợp lệ.')
   if d.get('confirm')!=target:return self.err(400,'Cần xác nhận chính xác tên tài khoản giáo viên.')
   if not throttle('reset:'+user['username'],12,3600):return self.err(429,'Đã thực hiện quá nhiều lượt đặt lại mật khẩu.')
   alphabet='ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789'
   temp=''.join((secrets.choice('ABCDEFGHJKLMNPQRSTUVWXYZ'),secrets.choice('abcdefghijkmnopqrstuvwxyz'),secrets.choice('23456789')))+''.join(secrets.choice(alphabet) for _ in range(17))
   context=TENANT.set(None)
   try:
    with db() as c:
     row=c.execute('SELECT role FROM users WHERE username=?',(target,)).fetchone()
     if not row or row['role']!='teacher':return self.err(404,'Không tìm thấy tài khoản giáo viên.')
     c.execute('UPDATE users SET password_hash=? WHERE username=?',(hashpw(temp),target))
     c.execute('INSERT INTO password_reset_required(username,issued) VALUES(?,?) ON CONFLICT(username) DO UPDATE SET issued=excluded.issued',(target,int(time.time())))
     c.execute('DELETE FROM login_sessions WHERE username=?',(target,))
   finally:TENANT.reset(context)
   return self.sendj(200,{'ok':True,'username':target,'temporaryPassword':temp,'message':'Mật khẩu tạm chỉ hiển thị một lần; hãy chuyển trực tiếp cho đúng giáo viên.'})
  if p=='/api/password/change':
   user=self.auth()
   if not user:return
   old=str(d.get('current_password',''))
   new=str(d.get('new_password',''))
   if not 12<=len(new)<=128 or not re.search('[A-Z]',new) or not re.search('[a-z]',new) or not re.search('[0-9]',new):
    return self.err(400,'Mật khẩu mới cần 12–128 ký tự, gồm chữ hoa, chữ thường và chữ số.')
   if old==new:return self.err(400,'Mật khẩu mới phải khác mật khẩu hiện tại.')
   context=TENANT.set(None)
   try:
    with db() as c:
     row=c.execute('SELECT password_hash FROM users WHERE username=?',(user['username'],)).fetchone()
     if not row or not checkpw(old,row['password_hash']):return self.err(401,'Mật khẩu hiện tại không chính xác.')
     c.execute('UPDATE users SET password_hash=? WHERE username=?',(hashpw(new),user['username']))
     c.execute('DELETE FROM password_reset_required WHERE username=?',(user['username'],))
     c.execute('DELETE FROM login_sessions WHERE username=?',(user['username'],))
   finally:TENANT.reset(context)
   return self.sendj(200,{'ok':True,'message':'Đổi mật khẩu thành công. Mọi phiên cũ đã hết hiệu lực; vui lòng đăng nhập lại.'})
  if p=='/api/games/answer':
   try:gid=int(d.get('game_id'));qid=int(d.get('question_id'));choice=str(d.get('choice','')).upper()
   except:return self.err(400,'Câu trả lời không hợp lệ.')
   if choice not in ('A','B','C','D'):return self.err(400,'Chọn A, B, C hoặc D.')
   with db() as c:
    game=c.execute('SELECT question_ids FROM games WHERE id=?',(gid,)).fetchone()
    if not game or qid not in json.loads(game['question_ids']):return self.err(404,'Câu hỏi không nằm trong trò chơi.')
    q=c.execute('SELECT answer,explanation FROM questions WHERE id=?',(qid,)).fetchone()
    if not q:return self.err(404,'Câu hỏi không tồn tại.')
   return self.sendj(200,{'correct':choice==q['answer'],'answer':q['answer'],'explanation':q['explanation']})
  if p=='/api/online/start':
   from online_exam import digest,make_form,valid_student
   code=str(d.get('exam_code','')).strip().upper()
   name=str(d.get('student_name','')).strip()
   student_code=str(d.get('student_code','')).strip().upper()
   if not valid_student(name,student_code):return self.err(400,'Cần họ tên (2–80 ký tự) và mã học sinh (3–32 ký tự chữ, số, gạch ngang).')
   if not re.fullmatch(r'[A-Z2-9]{12,24}',code):return self.err(400,'Mã dự thi không hợp lệ.')
   if not throttle('online-start:'+self.client_address[0],300,900):return self.err(429,'Quá nhiều lần yêu cầu bắt đầu bài thi.')
   context=TENANT.set(None)
   try:
    with db() as c:link=c.execute('SELECT teacher_username,exam_id FROM online_exam_links WHERE code_hash=? AND active=1',(digest(code),)).fetchone()
   finally:TENANT.reset(context)
   if not link:return self.err(404,'Không tìm thấy mã dự thi đang hoạt động.')
   teacher=link['teacher_username'];exam_id=link['exam_id']
   tenant=TENANT.set(tenant_for(teacher))
   try:
    with db() as c:
     exam=c.execute('SELECT id,title,subject,grade,duration,published FROM exams WHERE id=?',(exam_id,)).fetchone()
     if not exam or not exam['published']:return self.err(404,'Đề chưa được công bố.')
     if c.execute('SELECT 1 FROM online_exam_results WHERE exam_id=? AND student_code=?',(exam_id,student_code)).fetchone():
      return self.err(409,'Mã học sinh này đã nộp bài. Liên hệ giáo viên nếu cần kiểm tra lại.')
     questions=c.execute('SELECT q.* FROM exam_questions eq JOIN questions q ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY eq.ord',(exam_id,)).fetchall()
   finally:TENANT.reset(tenant)
   if not questions:return self.err(400,'Đề chưa có câu hỏi.')
   seed=secrets.randbits(63);token=secrets.token_urlsafe(32)
   expire=int(time.time())+int(exam['duration'])*60
   context=TENANT.set(None)
   try:
    with db() as c:
     c.execute('INSERT INTO online_exam_sessions(token_hash,code_hash,teacher_username,exam_id,student_name,student_code,seed,expires) VALUES(?,?,?,?,?,?,?,?)',(digest(token),digest(code),teacher,exam_id,name,student_code,seed,expire))
   finally:TENANT.reset(context)
   form=make_form(questions,seed)
   safe_form=[{'id':q['id'],'content':q['content'],'options':q['options']} for q in form]
   return self.sendj(200,{'attempt_token':token,'exam':{k:exam[k] for k in ('id','title','subject','grade','duration')},'expires':expire,'items':safe_form})
  if p=='/api/online/submit':
   from online_exam import digest,make_form,score_form
   token=str(d.get('attempt_token',''))
   answers=d.get('answers',{})
   if not re.fullmatch(r'[A-Za-z0-9_-]{30,128}',token) or not isinstance(answers,dict):return self.err(400,'Phiên làm bài hoặc câu trả lời không hợp lệ.')
   context=TENANT.set(None)
   try:
    with db() as c:run=c.execute('SELECT * FROM online_exam_sessions WHERE token_hash=?',(digest(token),)).fetchone()
    if run:
     with db() as c:link=c.execute('SELECT active FROM online_exam_links WHERE code_hash=?',(run['code_hash'],)).fetchone()
    else:link=None
   finally:TENANT.reset(context)
   if not run or not link:return self.err(404,'Phiên làm bài không tồn tại hoặc mã đã bị xóa.')
   # Re-issuing a teacher code blocks new starts, but must not discard exams already underway.
   if int(time.time())>int(run['expires'])+30:return self.err(403,'Đã hết thời gian nộp bài.')
   tenant=TENANT.set(tenant_for(run['teacher_username']))
   try:
    with db() as c:
     exam=c.execute('SELECT published FROM exams WHERE id=?',(run['exam_id'],)).fetchone()
     if not exam or not exam['published']:return self.err(404,'Đề đã đóng.')
     questions=c.execute('SELECT q.* FROM exam_questions eq JOIN questions q ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY eq.ord',(run['exam_id'],)).fetchall()
     forms=make_form(questions,int(run['seed']))
     outcome=score_form(forms,answers)
     score=outcome['score'];correct=outcome['correct'];total=len(forms)
     feedback=('Hoàn thành rất tốt.' if score>=8 else 'Đã đạt kiến thức cơ bản.' if score>=5 else 'Cần ôn tập thêm các nội dung chưa vững.')
     if c.execute('SELECT 1 FROM online_exam_results WHERE exam_id=? AND student_code=?',(run['exam_id'],run['student_code'])).fetchone():
      return self.err(409,'Mã học sinh đã có bài nộp cho đề này.')
     cur=c.execute('INSERT INTO attempts(exam_id,student_name,access_code_hash,form_seed,score,correct,total,answers,feedback) VALUES(?,?,?,?,?,?,?,?,?)',(run['exam_id'],run['student_name'],digest(run['code_hash']),run['seed'],score,correct,total,json.dumps(answers,ensure_ascii=False),feedback))
     c.execute('INSERT INTO online_exam_results(attempt_id,session_hash,exam_id,student_code) VALUES(?,?,?,?)',(cur.lastrowid,digest(token),run['exam_id'],run['student_code']))
   except sqlite3.IntegrityError:
    return self.err(409,'Bài thi đã được nộp hoặc mã học sinh đã có kết quả.')
   finally:TENANT.reset(tenant)
   context=TENANT.set(None)
   try:
    with db() as c:c.execute('DELETE FROM online_exam_sessions WHERE token_hash=?',(digest(token),))
   finally:TENANT.reset(context)
   return self.sendj(200,{'score':score,'correct':correct,'total':total,'feedback':feedback,'message':'Đã lưu bài làm và tự động chấm điểm.'})
  if p=='/api/exams/submit':
   try:eid=int(d.get('exam_id')); name=str(d.get('student_name','')).strip()[:80];code=str(d.get('access_code','')).strip();answers=d.get('answers',{})
   except:return self.err(400,'Thông tin bài làm không hợp lệ.')
   if not name or len(name)<2 or not isinstance(answers,dict) or len(code)<4:return self.err(400,'Nhập tên và mã dự thi ít nhất 4 ký tự.')
   with db() as c:
    e=c.execute('SELECT * FROM exams WHERE id=? AND published=1',(eid,)).fetchone()
    if not e:return self.err(404,'Đề chưa công bố.')
    qs=c.execute('SELECT q.* FROM exam_questions eq JOIN questions q ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY eq.ord',(eid,)).fetchall()
    seed=int.from_bytes(hmac.new(SECRET.encode(),f'{eid}:{name}:{code}'.encode(),hashlib.sha256).digest()[:8],'big') if SECRET else int.from_bytes(hashlib.sha256(f'{eid}:{name}:{code}'.encode()).digest()[:8],'big')
    seed = seed & ((1 << 63) - 1)
    forms=shuffled(qs,seed);correct=0;report=[]
    for q in forms:
     selected=str(answers.get(str(q['id']),'')).upper();ok=selected==q['_correct'];correct+=ok
     orig=next(x for x in qs if x['id']==q['id'])
     report.append({'question':q['content'],'selected':selected,'correct':q['_correct'],'ok':ok,'explanation':orig['explanation']})
    total=len(forms);score=round(correct/total*10,2) if total else 0
    fb=('Hoàn thành rất tốt. ' if score>=8 else 'Đã nắm được phần cơ bản. ' if score>=5 else 'Nên ôn tập lại các khái niệm nền tảng. ')+f'Đúng {correct}/{total} câu. Hãy xem giải thích ở các câu chưa đúng.'
    c.execute('INSERT INTO attempts(exam_id,student_name,access_code_hash,form_seed,score,correct,total,answers,feedback) VALUES(?,?,?,?,?,?,?,?,?)',(eid,name,hashlib.sha256((code+name).encode()).hexdigest(),seed,score,correct,total,json.dumps(answers),fb))
   return self.sendj(200,{'score':score,'correct':correct,'total':total,'feedback':fb,'review':report,'ai':False})
  if p=='/api/exams/form':
   try:eid=int(d.get('exam_id'));name=str(d.get('student_name','')).strip()[:80];code=str(d.get('access_code','')).strip()
   except:return self.err(400,'Dữ liệu không hợp lệ.')
   if len(name)<2 or len(code)<4:return self.err(400,'Cần họ tên và mã dự thi từ 4 ký tự.')
   with db() as c:
    e=c.execute('SELECT id,title,subject,grade,duration FROM exams WHERE id=? AND published=1',(eid,)).fetchone()
    if not e:return self.err(404,'Không tìm thấy đề công bố.')
    qs=c.execute('SELECT q.* FROM exam_questions eq JOIN questions q ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY eq.ord',(eid,)).fetchall()
   seed=int.from_bytes(hmac.new(SECRET.encode(),f'{eid}:{name}:{code}'.encode(),hashlib.sha256).digest()[:8],'big') if SECRET else int.from_bytes(hashlib.sha256(f'{eid}:{name}:{code}'.encode()).digest()[:8],'big')
   seed = seed & ((1 << 63) - 1)
   form=shuffled(qs,seed)
   return self.sendj(200,{'exam':dict(e),'items':[{'id':q['id'],'content':q['content'],'options':q['options']} for q in form]})
  s=self.auth()
  if not s:return
  if p=='/api/online/share':
   from online_exam import digest,create_exam_code
   try:exam_id=int(d.get('exam_id'))
   except:return self.err(400,'Mã đề không hợp lệ.')
   with db() as c:
    e=c.execute('SELECT id,published FROM exams WHERE id=?',(exam_id,)).fetchone()
    if not e or not e['published']:return self.err(404,'Hãy công bố đề trước khi cấp mã dự thi.')
    if c.execute('SELECT COUNT(*) FROM exam_questions WHERE exam_id=?',(exam_id,)).fetchone()[0]<1:return self.err(400,'Đề chưa có câu hỏi.')
   code=create_exam_code()
   context=TENANT.set(None)
   try:
    with db() as c:
     c.execute('UPDATE online_exam_links SET active=0 WHERE teacher_username=? AND exam_id=?',(s['username'],exam_id))
     c.execute('INSERT INTO online_exam_links(code_hash,teacher_username,exam_id) VALUES(?,?,?)',(digest(code),s['username'],exam_id))
   finally:TENANT.reset(context)
   return self.sendj(200,{'exam_id':exam_id,'exam_code':code,'url':PUBLIC_URL.rstrip('/')+'/student' if PUBLIC_URL else '/student','message':'Mã mới chỉ hiện một lần. Cấp riêng cho học sinh; mã cũ đã hết hiệu lực.'})
  if p=='/api/exams/word/publish':
   from exam_word import validate_questions
   try:
    questions=validate_questions(d.get('items'))
    title=str(d.get('title','')).strip()
    duration=int(d.get('duration',45))
    subject=str(d.get('subject','Tin học')).strip()[:80]
    grade=str(d.get('grade','9')).strip()[:20]
    if not 3<=len(title)<=130 or not 1<=duration<=240 or len(questions)>100:raise ValueError('Tên đề, thời lượng hoặc số câu không hợp lệ (tối đa 100 câu).')
   except (ValueError,TypeError) as exc:return self.err(400,str(exc))
   with db() as c:
    ex=c.execute('INSERT INTO exams(title,subject,grade,duration,published) VALUES(?,?,?,?,?)',(title,subject,grade,duration,1))
    eid=ex.lastrowid
    for i,q in enumerate(questions):
     option=q['options']
     cur=c.execute('INSERT INTO questions(subject,grade,topic,content,a,b,c,d,answer,origin) VALUES(?,?,?,?,?,?,?,?,?,?)',(subject,grade,'Nhập từ Word',q['content'],option['A'],option['B'],option['C'],option['D'],q['answer'],'word-import'))
     c.execute('INSERT INTO exam_questions(exam_id,question_id,ord) VALUES(?,?,?)',(eid,cur.lastrowid,i))
   return self.sendj(201,{'id':eid,'count':len(questions),'message':'Đã lưu đề Word vào hệ thống và công bố để học sinh làm trực tuyến.'})
  if p=='/api/exams/word/preview':
   from exam_word import parse_word
   try:
    filename=str(d.get('filename','')).strip()
    if not filename.lower().endswith('.docx'):return self.err(400,'Vui lòng chọn file Word .docx (không hỗ trợ .doc).')
    encoded=str(d.get('file',''))
    if len(encoded)>2_800_000:return self.err(400,'File Word tối đa 2 MB.')
    raw=base64.b64decode(encoded,validate=True)
    items=parse_word(raw)
    return self.sendj(200,{'items':items,'count':len(items),'message':'Hãy kiểm tra nội dung từng câu và đáp án trước khi trộn.'})
   except (ValueError,TypeError,base64.binascii.Error) as exc:return self.err(400,str(exc)[:350])
   except Exception:return self.err(400,'Không đọc được tài liệu Word. Hãy lưu lại dưới dạng .docx hợp lệ.')
  if p=='/api/exams/word/export':
   from exam_word import export_package
   try:
    count=int(d.get('variants',0))
    title=str(d.get('title','')).strip()
    data=export_package(title,d.get('items'),count)
   except (ValueError,TypeError) as exc:return self.err(400,str(exc)[:350])
   if len(data)>15_000_000:return self.err(413,'Đề tạo ra quá lớn để xuất trong một lần.')
   self.send_response(200)
   self.send_header('Content-Type','application/zip')
   self.send_header('Content-Disposition','attachment; filename="De_tron_va_dap_an.zip"')
   self.send_header('Content-Length',str(len(data)))
   self.send_header('Cache-Control','private, no-store')
   self.send_header('X-Content-Type-Options','nosniff')
   self.end_headers();return self.wfile.write(data)
  if p=='/api/ai/chat':
   if not AI_KEY:return self.err(503,'Chưa cấu hình Gemini API trên máy chủ.')
   if not throttle('chat:'+s['username'],12,60):return self.err(429,'Bạn đã gửi nhiều tin nhắn. Vui lòng thử lại sau một phút.')
   prompt=str(d.get('message','')).strip()[:1800]
   if not prompt:return self.err(400,'Hãy nhập nội dung muốn hỏi.')
   raw_history=d.get('history',[])
   if not isinstance(raw_history,list):return self.err(400,'Lịch sử hội thoại không hợp lệ.')
   history=[]
   for turn in raw_history[-8:]:
    if isinstance(turn,dict) and turn.get('role') in ('user','assistant') and isinstance(turn.get('text'),str):
     history.append((turn['role'],turn['text'][:900]))
   context='\n'.join(('Giáo viên: ' if role=='user' else 'Trợ lý: ')+text for role,text in history)
   instruction=('Bạn là trợ lý giáo dục SmartClass AI, trả lời bằng tiếng Việt rõ ràng, lịch sự; hỗ trợ giáo viên các môn học, gợi ý bài dạy theo CV5512, xây dựng câu hỏi và hướng dẫn sử dụng ứng dụng. Nếu chưa chắc chắn hãy nói rõ. Không yêu cầu hoặc tiết lộ API Key, mật khẩu hay thông tin riêng của học sinh. Không tự nhận kết quả AI đã được kiểm chứng.\n'+context+'\nGiáo viên: '+prompt+'\nTrợ lý:')
   try:
    result=gemini([{'text':instruction}],timeout=65)
    answer=''.join(part.get('text','') for cand in result.get('candidates',[]) for part in cand.get('content',{}).get('parts',[]) if isinstance(part,dict)).strip()
    if not answer:return self.err(502,'Gemini chưa trả nội dung. Hãy thử lại.')
    return self.sendj(200,{'reply':answer[:8500],'model':AI_MODEL})
   except (RuntimeError,ValueError,KeyError,TypeError) as exc:
    return self.err(502,str(exc) if isinstance(exc,RuntimeError) else 'Phản hồi Gemini chưa hợp lệ.')
  if p=='/api/logout':
   cookie=SimpleCookie();cookie.load(self.headers.get('Cookie',''))
   if 'scsession' in cookie:
    context=TENANT.set(None)
    try:
     with db() as c:
      c.execute('DELETE FROM login_sessions WHERE token_hash=?',(hashlib.sha256(cookie['scsession'].value.encode()).hexdigest(),))
    finally:TENANT.reset(context)
   return self.sendj(200,{'ok':True},{'Set-Cookie':'scsession=; Max-Age=0; HttpOnly; SameSite=Lax; Path=/'})
  if p=='/api/planner/ai-draft':
   # Teacher-only; the do_POST authentication/CSRF guard runs before this block.
   if not AI_KEY:return self.err(503,'Chưa cấu hình GEMINI_API_KEY. Bạn vẫn có thể soạn thủ công và xuất Word.')
   if not throttle('planner-ai:'+s['username'],4,60):return self.err(429,'Đã gửi quá nhiều yêu cầu soạn giáo án; thử lại sau một phút.')
   keys=('subject','grade','title','periods','knowledge','skills','qualities','digital','digital_evidence','resources','warmup','explore','practice','apply')
   values={k:str(d.get(k,'')).strip()[:1100] for k in keys}
   if len(values['title'])<4 or not values['subject'] or not values['grade']:return self.err(400,'Nhập môn học, lớp và tên bài dạy trước khi nhờ AI soạn.')
   try:
    periods=int(values['periods'] or '1')
    if not 1<=periods<=20:raise ValueError()
   except:return self.err(400,'Thời lượng bài dạy cần từ 1 đến 20 tiết.')
   instruction=('Bạn là trợ lý soạn KẾ HOẠCH BÀI DẠY theo PHỤ LỤC IV Công văn 5512/BGDĐT-GDTrH. '
     'Soạn dự thảo bằng tiếng Việt, không tự tuyên bố đạt chuẩn; chỉ dùng yêu cầu cần đạt được giáo viên cung cấp, nếu chưa có thì nêu cần đối chiếu CTGDPT và SGK. '
     'Phát triển năng lực số tham khảo Thông tư 02/2025/TT-BGDĐT, phải chọn minh chứng khả thi, KHÔNG ép sử dụng thiết bị cho tất cả học sinh. '
     'Mỗi hoạt động cần có bốn thành phần ghi rõ: Mục tiêu; Nội dung; Sản phẩm học tập; Tổ chức thực hiện (giao nhiệm vụ, thực hiện, báo cáo/thảo luận, kết luận/đánh giá). '
     'Chú ý đánh giá quá trình và phân hóa nhiệm vụ; kiến thức phải đúng môn và lớp. Không đưa thông tin học sinh thật. '
     'Chỉ trả JSON object chứa ĐỦ các khóa knowledge,skills,qualities,digital_evidence,resources,warmup,explore,practice,apply. '
     'Mỗi giá trị là chuỗi dễ đọc, tuyệt đối không Markdown fenced code. '
     'Dữ liệu giáo viên nhập: '+json.dumps(values,ensure_ascii=False))
   try:
    result=gemini_text(instruction)
    if not isinstance(result,dict):raise ValueError('Gemini không trả một đối tượng giáo án.')
    required=('knowledge','skills','qualities','digital_evidence','resources','warmup','explore','practice','apply')
    fields={k:str(result.get(k,'')).strip()[:6500] for k in required}
    if any(len(fields[k])<18 for k in required):raise ValueError('Gemini trả thiếu nội dung của kế hoạch; hãy thử lại.')
    return self.sendj(200,{'draft':fields,'note':'Bản nháp AI. Giáo viên cần đối chiếu yêu cầu cần đạt, điều kiện lớp học và rà soát trước khi xuất Word.'})
   except (RuntimeError,ValueError,KeyError,TypeError) as exc:
    return self.err(502,str(exc)[:400] if isinstance(exc,(RuntimeError,ValueError)) else 'Phản hồi Gemini không hợp lệ.')
  if p=='/api/planner/docx':
   from docx import Document
   from docx.shared import Cm, Pt
   from docx.enum.text import WD_ALIGN_PARAGRAPH
   textfields=['school','teacher','subject','grade','title','periods','knowledge','skills','qualities','digital','digital_evidence','resources','warmup','explore','practice','apply']
   vals={k:str(d.get(k,'')).strip()[:8000] for k in textfields}
   if not vals['title'] or not vals['subject']:return self.err(400,'Vui lòng nhập môn học và tên bài dạy.')
   doc=Document();sec=doc.sections[0];sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
   normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(13)
   top=doc.add_paragraph(f"Trường: {vals['school']}\nGiáo viên: {vals['teacher']}")
   heading=doc.add_heading('KẾ HOẠCH BÀI DẠY',0);heading.alignment=WD_ALIGN_PARAGRAPH.CENTER
   doc.add_heading(vals['title'],level=1)
   doc.add_paragraph(f"Môn học: {vals['subject']}     Lớp: {vals['grade']}     Thời gian thực hiện: {vals['periods']} tiết")
   doc.add_heading('I. MỤC TIÊU',level=1)
   for title,key in [('1. Kiến thức','knowledge'),('2. Năng lực','skills'),('3. Phẩm chất','qualities'),('4. Năng lực số dự kiến phát triển','digital'),('Minh chứng đánh giá năng lực số','digital_evidence')]:
    doc.add_heading(title,level=2);doc.add_paragraph(vals[key] or '(Giáo viên bổ sung)')
   doc.add_heading('II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU',level=1);doc.add_paragraph(vals['resources'])
   doc.add_heading('III. TIẾN TRÌNH DẠY HỌC',level=1)
   for title,key in [('Hoạt động 1. Mở đầu','warmup'),('Hoạt động 2. Hình thành kiến thức','explore'),('Hoạt động 3. Luyện tập','practice'),('Hoạt động 4. Vận dụng','apply')]:
    doc.add_heading(title,level=2);doc.add_paragraph(vals[key] or '(Giáo viên bổ sung mục tiêu, nội dung, sản phẩm, tổ chức thực hiện)')
   doc.add_paragraph('Lưu ý: Mẫu hỗ trợ soạn theo Phụ lục IV CV5512; giáo viên tự rà soát nội dung và yêu cầu chuyên môn.')
   buffer=io.BytesIO();doc.save(buffer);raw=buffer.getvalue();self.send_response(200)
   self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
   self.send_header('Content-Disposition','attachment; filename="Ke_hoach_bai_day_5512.docx"')
   self.send_header('Content-Length',str(len(raw)));self.send_header('Cache-Control','no-store');self.end_headers();return self.wfile.write(raw)
  if p=='/api/classes':
   name=str(d.get('name','')).strip()[:100];grade=str(d.get('grade','')).strip()[:25]
   if len(name)<2:return self.err(400,'Tên lớp cần ít nhất 2 ký tự.')
   try:
    with db() as c:cur=c.execute('INSERT INTO classes(name,grade) VALUES(?,?)',(name,grade));idx=cur.lastrowid
    return self.sendj(201,{'id':idx})
   except sqlite3.IntegrityError:return self.err(409,'Lớp đã tồn tại.')
  if p=='/api/students':
   name=str(d.get('name','')).strip()[:100];code=str(d.get('code','')).strip()[:40]
   try:cid=int(d.get('class_id',0))
   except:return self.err(400,'Chọn lớp hợp lệ.')
   if len(name)<2 or not re.fullmatch(r'[A-Za-z0-9_.-]{2,40}',code):return self.err(400,'Tên hoặc mã học sinh không hợp lệ (mã chỉ dùng chữ, số, dấu -_.).')
   try:
    with db() as c:
     if not c.execute('SELECT 1 FROM classes WHERE id=?',(cid,)).fetchone():return self.err(404,'Không tìm thấy lớp.')
     cur=c.execute('INSERT INTO students(code,name,class_id) VALUES(?,?,?)',(code,name,cid));idx=cur.lastrowid
    return self.sendj(201,{'id':idx})
   except sqlite3.IntegrityError:return self.err(409,'Mã học sinh đã tồn tại.')
  if p=='/api/students/delete':
   try:sid=int(d.get('id'))
   except:return self.err(400,'Mã học sinh không hợp lệ.')
   with db() as c:
    c.execute('DELETE FROM attendance WHERE student_id=?',(sid,));cur=c.execute('DELETE FROM students WHERE id=?',(sid,))
   return self.sendj(200,{'deleted':cur.rowcount})
  if p=='/api/attendance':
   day=str(d.get('day',''));rows=d.get('items',[])
   if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',day) or not isinstance(rows,list) or len(rows)>150:return self.err(400,'Dữ liệu điểm danh không hợp lệ.')
   try:datetime.strptime(day,'%Y-%m-%d')
   except:return self.err(400,'Ngày không hợp lệ.')
   allowed=('present','absent','late','excused')
   parsed=[]
   for r in rows:
    try:sid=int(r['student_id']);status=str(r['status']);note=str(r.get('note',''))[:200]
    except:return self.err(400,'Dữ liệu học sinh không hợp lệ.')
    if status not in allowed:return self.err(400,'Trạng thái điểm danh không hợp lệ.')
    parsed.append((sid,day,status,note))
   with db() as c:
    found=c.execute('SELECT COUNT(*) FROM students WHERE id IN ('+','.join('?' for _ in parsed)+')',[a[0] for a in parsed]).fetchone()[0] if parsed else 0
    if found!=len(set(a[0] for a in parsed)) or found!=len(parsed):return self.err(400,'Có mã học sinh không hợp lệ hoặc trùng.')
    # Save 50 students in each SQL statement rather than N network round trips.
    # All batches share one transaction, so failure cannot partially save a class.
    for offset in range(0,len(parsed),50):
     batch=parsed[offset:offset+50]
     values_sql=','.join('(?,?,?,?)' for _ in batch)
     params=[item for record in batch for item in record]
     c.execute('INSERT INTO attendance(student_id,day,status,note) VALUES '+values_sql+' ON CONFLICT(student_id,day) DO UPDATE SET status=excluded.status,note=excluded.note',params)
    if parsed:
     ids=[row[0] for row in parsed];class_ids=set()
     for offset in range(0,len(ids),300):
      chunk=ids[offset:offset+300]
      for row in c.execute('SELECT DISTINCT class_id FROM students WHERE id IN ('+','.join('?' for _ in chunk)+')',chunk):
       if row['class_id'] is not None:class_ids.add(row['class_id'])
     for cid in class_ids:
      c.execute('INSERT INTO attendance_days(class_id,day) VALUES(?,?) ON CONFLICT(class_id,day) DO NOTHING',(cid,day))
   return self.sendj(200,{'saved':len(parsed),'day':day})
  if p=='/api/classes/rename':
   try:cid=int(d.get('id',0))
   except:return self.err(400,'Mã lớp không hợp lệ.')
   name=str(d.get('name','')).strip()[:100]
   grade=str(d.get('grade','')).strip()[:25]
   if cid<=0 or len(name)<2:return self.err(400,'Tên lớp phải có ít nhất 2 ký tự.')
   try:
    with db() as c:
     row=c.execute('SELECT id FROM classes WHERE id=?',(cid,)).fetchone()
     if not row:return self.err(404,'Không tìm thấy lớp để sửa.')
     c.execute('UPDATE classes SET name=?,grade=? WHERE id=?',(name,grade,cid))
    return self.sendj(200,{'ok':True,'id':cid,'name':name,'grade':grade})
   except sqlite3.IntegrityError:return self.err(409,'Tên lớp đã tồn tại, hãy chọn tên khác.')
  if p=='/api/classes/delete':
   try:cid=int(d.get('id'))
   except:return self.err(400,'Mã lớp không hợp lệ.')
   with db() as c:
    cl=c.execute('SELECT name FROM classes WHERE id=?',(cid,)).fetchone()
    if not cl:return self.err(404,'Lớp không tồn tại.')
    count=c.execute('SELECT COUNT(*) FROM students WHERE class_id=?',(cid,)).fetchone()[0]
    if count:return self.err(409,'Lớp còn '+str(count)+' học sinh. Hãy chuyển học sinh sang lớp khác trước khi xóa.')
    c.execute('DELETE FROM classes WHERE id=?',(cid,))
   return self.sendj(200,{'deleted':1})
  if p=='/api/students/move':
   try:ids=list(dict.fromkeys(int(x) for x in d.get('ids',[])));cid=int(d.get('class_id'))
   except:return self.err(400,'Danh sách học sinh hoặc lớp không hợp lệ.')
   if not ids or len(ids)>2000 or any(i<=0 for i in ids):return self.err(400,'Chọn từ 1–2000 học sinh.')
   with db() as c:
    if not c.execute('SELECT 1 FROM classes WHERE id=?',(cid,)).fetchone():return self.err(404,'Không tìm thấy lớp đích.')
    count=c.execute('SELECT COUNT(*) FROM students WHERE id IN ('+','.join('?' for _ in ids)+')',ids).fetchone()[0]
    if count!=len(ids):return self.err(400,'Một số học sinh không tồn tại.')
    for offset in range(0,len(ids),350):
     chunk=ids[offset:offset+350]
     c.execute('UPDATE students SET class_id=? WHERE id IN ('+','.join('?' for _ in chunk)+')',[cid]+chunk)
   return self.sendj(200,{'moved':len(ids)})
  if p=='/api/students/bulk-delete':
   try:ids=list(dict.fromkeys(int(x) for x in d.get('ids',[])))
   except:return self.err(400,'Danh sách học sinh không hợp lệ.')
   if not ids or len(ids)>2000 or d.get('confirm')!='XOA_HOC_SINH':return self.err(400,'Cần xác nhận xóa học sinh.')
   with db() as c:
    count=c.execute('SELECT COUNT(*) FROM students WHERE id IN ('+','.join('?' for _ in ids)+')',ids).fetchone()[0]
    if count!=len(ids):return self.err(400,'Một số học sinh không tồn tại.')
    c.executemany('DELETE FROM students WHERE id=?',[(i,) for i in ids])
   return self.sendj(200,{'deleted':len(ids),'warning':'Dữ liệu điểm danh liên quan đã bị xóa.'})
  if p=='/api/students/import':
   try:
    raw=base64.b64decode(str(d.get('file','')),validate=True)
    if len(raw)>2500000:return self.err(400,'Tệp Excel tối đa 2,5 MB.')
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    ws=wb.active; rows=ws.iter_rows(values_only=True)
    head=next(rows,None)
    if not head:raise ValueError('Excel không có tiêu đề')
    def normalize(x):
     import unicodedata
     return ''.join(ch for ch in unicodedata.normalize('NFD',str(x or '').strip().lower()) if unicodedata.category(ch)!='Mn').replace('đ','d').replace('_','').replace(' ','')
    headers=[normalize(x) for x in head]
    def find(keys):return next((i for i,h in enumerate(headers) if h in keys),None)
    ci=find(('mahocsinh','mahs','mahocsinhid','code','studentid','mssv'))
    ni=find(('hovaten','hoten','tenhocsinh','fullname','name'))
    li=find(('lop','tenlop','class','classname'))
    if ci is None or ni is None or li is None:raise ValueError('Cần đủ 3 cột: Mã học sinh, Họ và tên, Lớp.')
    entries=[];seen=set()
    for lineno,row in enumerate(rows,2):
     if lineno>2002:raise ValueError('Tối đa 2000 học sinh/lần.')
     vals=[str(row[i] or '').strip() if i<len(row) else '' for i in (ci,ni,li)]
     if not any(vals):continue
     code,name,cl=vals
     if not re.fullmatch(r'[A-Za-z0-9_.-]{2,40}',code) or not 2<=len(name)<=100 or not 1<=len(cl)<=100:raise ValueError('Dòng '+str(lineno)+' thiếu hoặc sai mã, họ tên, lớp.')
     if code in seen:raise ValueError('Mã học sinh trùng trong Excel tại dòng '+str(lineno)+': '+code)
     seen.add(code);entries.append((code,name,cl))
    wb.close()
    if not entries:return self.err(400,'Không tìm thấy học sinh trong Excel.')
    from collections import Counter
    summary=dict(sorted(Counter(x[2] for x in entries).items()))
    if d.get('preview') is True:return self.sendj(200,{'preview':True,'total':len(entries),'classes':summary})
    if d.get('confirm')!='NHAP_TOAN_TRUONG':return self.err(400,'Cần xem trước phân lớp và xác nhận nhập.')
    # Batch SQL calls rather than querying/inserting each student individually.
    with db() as c:
     existing_classes={r['name'] for r in c.execute('SELECT name FROM classes')}
     new_classes=[(cl,''.join(ch for ch in cl if ch.isdigit())[:2]) for cl in summary if cl not in existing_classes]
     if new_classes:c.executemany('INSERT OR IGNORE INTO classes(name,grade) VALUES(?,?)',new_classes)
     classids={r['name']:r['id'] for r in c.execute('SELECT id,name FROM classes')}
     codes=[item[0] for item in entries]
     existing_codes=set()
     for offset in range(0,len(codes),350):
      chunk=codes[offset:offset+350]
      existing_codes.update(row['code'] for row in c.execute('SELECT code FROM students WHERE code IN ('+','.join('?' for _ in chunk)+')',chunk))
     created=sum(code not in existing_codes for code in codes)
     updated=len(entries)-created
     payload=[(code,name,classids[cl]) for code,name,cl in entries]
     for offset in range(0,len(payload),250):
      c.executemany('INSERT INTO students(code,name,class_id) VALUES(?,?,?) ON CONFLICT(code) DO UPDATE SET name=excluded.name,class_id=excluded.class_id',payload[offset:offset+250])
     classes_created=len(new_classes)
    return self.sendj(200,{'imported':len(entries),'created':created,'updated':updated,'classes_created':classes_created,'classes':summary})
   except Exception as e:return self.err(400,'Không đọc được Excel: '+str(e)[:180])
  if p=='/api/demo/multisubject':
   # Teacher-only demo fixture. Never pass demo scores off as genuine learning outcomes.
   sample=[('Toán', '8', 'Đại số', 'Giá trị của 3² là bao nhiêu?', '6', '9', '12', '8', 'B', 'Nhận biết', '3² = 3 × 3 = 9.', 'demo-multisubject'), ('Ngữ văn', '8', 'Tiếng Việt', 'Từ nào sau đây là từ láy?', 'Học sinh', 'Lấp lánh', 'Sách vở', 'Bàn ghế', 'B', 'Nhận biết', 'Lấp lánh là từ láy.', 'demo-multisubject'), ('Tiếng Anh', '8', 'Ngữ pháp', 'Choose the correct form: She ___ to school every day.', 'go', 'goes', 'going', 'gone', 'B', 'Nhận biết', 'She goes: ngôi thứ ba số ít ở hiện tại đơn.', 'demo-multisubject'), ('Vật lí', '8', 'Chuyển động', 'Đơn vị SI của vận tốc là gì?', 'm/s', 'kg', 'N', 'J', 'A', 'Nhận biết', 'Vận tốc có đơn vị SI là mét trên giây.', 'demo-multisubject'), ('Hóa học', '8', 'Chất và nguyên tử', 'Kí hiệu hóa học của nguyên tố oxi là gì?', 'Ox', 'O', 'Om', 'H', 'B', 'Nhận biết', 'O là kí hiệu nguyên tố oxi.', 'demo-multisubject'), ('Sinh học', '8', 'Tế bào', 'Bộ phận nào giúp tế bào thực vật quang hợp?', 'Nhân', 'Lục lạp', 'Màng sinh chất', 'Không bào', 'B', 'Nhận biết', 'Lục lạp có diệp lục thực hiện quang hợp.', 'demo-multisubject'), ('Lịch sử', '8', 'Lịch sử Việt Nam', 'Chiến thắng Bạch Đằng năm 938 gắn với ai?', 'Ngô Quyền', 'Trần Hưng Đạo', 'Lê Lợi', 'Quang Trung', 'A', 'Nhận biết', 'Ngô Quyền lãnh đạo chiến thắng Bạch Đằng năm 938.', 'demo-multisubject'), ('Địa lí', '8', 'Địa lí tự nhiên', 'Lớp khí bao quanh Trái Đất được gọi là gì?', 'Thủy quyển', 'Khí quyển', 'Thạch quyển', 'Sinh quyển', 'B', 'Nhận biết', 'Khí quyển là lớp khí bao quanh Trái Đất.', 'demo-multisubject'), ('Giáo dục công dân', '8', 'Quyền và nghĩa vụ', 'Hành vi nào thể hiện tôn trọng người khác?', 'Lắng nghe ý kiến', 'Chế giễu bạn', 'Ngắt lời', 'Lan truyền tin đồn', 'A', 'Nhận biết', 'Lắng nghe là biểu hiện tôn trọng.', 'demo-multisubject'), ('Tin học', '8', 'An toàn số', 'Mật khẩu nào an toàn hơn?', '123456', 'password', 'Một mật khẩu dài và riêng biệt', 'Ngày sinh', 'C', 'Nhận biết', 'Mật khẩu dài, khó đoán và riêng cho từng tài khoản an toàn hơn.', 'demo-multisubject'), ('Công nghệ', '8', 'Thiết kế kĩ thuật', 'Bước đầu của quy trình thiết kế kĩ thuật là gì?', 'Chế tạo hàng loạt', 'Xác định vấn đề', 'Đóng gói sản phẩm', 'Quảng cáo', 'B', 'Nhận biết', 'Cần xác định vấn đề trước khi đề xuất giải pháp.', 'demo-multisubject'), ('Khoa học tự nhiên', '8', 'Kĩ năng thực nghiệm', 'Dụng cụ nào thường dùng để đo nhiệt độ?', 'Thước', 'Cân', 'Nhiệt kế', 'Đồng hồ', 'C', 'Nhận biết', 'Nhiệt kế dùng để đo nhiệt độ.', 'demo-multisubject'), ('Âm nhạc', '8', 'Nhạc lí', 'Một nốt đen thông thường có trường độ bằng mấy nốt móc đơn?', '1', '2', '3', '4', 'B', 'Nhận biết', 'Một nốt đen bằng hai nốt móc đơn.', 'demo-multisubject'), ('Mĩ thuật', '8', 'Màu sắc', 'Ba màu cơ bản truyền thống trong hội họa là gì?', 'Đỏ, vàng, xanh lam', 'Đen, trắng, xám', 'Lục, tím, nâu', 'Cam, lục, tím', 'A', 'Nhận biết', 'Bộ màu cơ bản truyền thống: đỏ, vàng, xanh lam.', 'demo-multisubject'), ('Giáo dục thể chất', '8', 'An toàn luyện tập', 'Trước khi vận động mạnh nên làm gì?', 'Bỏ khởi động', 'Khởi động phù hợp', 'Uống thật nhiều nước một lúc', 'Tập ngay tối đa', 'B', 'Nhận biết', 'Khởi động giúp chuẩn bị cơ thể cho vận động.', 'demo-multisubject'), ('Hoạt động trải nghiệm', '8', 'Kĩ năng hợp tác', 'Khi làm việc nhóm, cách nào giúp hợp tác hiệu quả?', 'Không chia sẻ nhiệm vụ', 'Phân công rõ và lắng nghe', 'Chỉ nghe một người', 'Không phản hồi', 'B', 'Nhận biết', 'Phân công và lắng nghe giúp nhóm hợp tác.', 'demo-multisubject')]
   made=0
   with db() as c:
    for row in sample:
     exists=c.execute('SELECT id FROM questions WHERE subject=? AND content=? AND origin=?',(row[0],row[3],'demo-multisubject')).fetchone()
     if exists:continue
     c.execute('INSERT INTO questions(subject,grade,topic,content,a,b,c,d,answer,difficulty,explanation,origin) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',row)
     made+=1
    for subject in SUBJECT_CATALOG:
     qids=[v[0] for v in c.execute('SELECT id FROM questions WHERE subject=? AND origin=? ORDER BY id LIMIT 10',(subject,'demo-multisubject')).fetchall()]
     if not qids:continue
     title='Ôn tập '+subject+' [MINH HỌA]'
     if c.execute('SELECT 1 FROM exams WHERE title=?',(title,)).fetchone():continue
     eid=c.execute('INSERT INTO exams(title,subject,grade,duration,published) VALUES(?,?,?,?,0)',(title,subject,'8',15)).lastrowid
     c.executemany('INSERT INTO exam_questions(exam_id,question_id,ord) VALUES(?,?,?)',[(eid,q,i) for i,q in enumerate(qids)])
   return self.sendj(200,{'ok':True,'added':made,'subjects':len(SUBJECT_CATALOG),'notice':'DỮ LIỆU MINH HỌA: Giáo viên phải rà soát đáp án và sửa theo chương trình từng lớp. Đề mẫu mặc định chưa công bố.'})
  if p=='/api/demo/showcase':
   # Teacher-authenticated, reversible-by-selecting-demo-data fixture. No fabricated student outcomes.
   added={'games':0,'media':0,'attendance':0,'draft_exams':0}
   with db() as c:
    subjects=[r[0] for r in c.execute("SELECT DISTINCT subject FROM questions WHERE origin='demo-multisubject' ORDER BY subject")]
    for subject in subjects:
     ids=[r[0] for r in c.execute("SELECT id FROM questions WHERE subject=? AND origin='demo-multisubject' ORDER BY id LIMIT 10",(subject,))]
     if not ids:continue
     for kind,label in [('quiz','Trắc nghiệm'),('flash','Lật thẻ'),('speed','Đua tốc độ'),('mystery','Ô cửa bí mật'),('ladder','Leo thang')]:
      name=f'{label} · {subject} [MINH HỌA]'
      if not c.execute('SELECT id FROM games WHERE title=?',(name,)).fetchone():
       c.execute('INSERT INTO games(title,topic,kind,question_ids) VALUES(?,?,?,?)',(name,subject,kind,json.dumps(ids)))
       added['games']+=1
    lesson_title='Kế hoạch bài dạy mẫu · Toán 8 [MINH HỌA]'
    planner={ 'school':'TRƯỜNG THCS MINH HỌA', 'teacher':'GIÁO VIÊN MINH HỌA', 'subject':'Toán', 'grade':'8', 'periods':'1', 'title':'Hằng đẳng thức đáng nhớ (dữ liệu minh họa)', 'knowledge':'Nhận biết và vận dụng công thức bình phương một tổng.', 'skills':'Tư duy và lập luận toán học; giải quyết vấn đề.', 'qualities':'Chăm chỉ, trung thực trong trình bày lời giải.', 'digital':'Khai thác thông tin và dữ liệu; sáng tạo nội dung số', 'digital_evidence':'Sử dụng bảng tính hoặc công cụ trực tuyến để kiểm tra các ví dụ.', 'resources':'Máy chiếu; phiếu học tập; công cụ số phù hợp.', 'warmup':'Nêu tình huống tính nhanh (a+b)²; học sinh dự đoán kết quả.', 'explore':'Hướng dẫn hình thành công thức và giải thích bằng mô hình diện tích.', 'practice':'Học sinh thực hiện 3 bài tập có mức độ tăng dần, đối chiếu và sửa lỗi.', 'apply':'Vận dụng công thức giải một bài toán thực tiễn; ghi nhận sản phẩm.' }
    if not c.execute('SELECT id FROM lesson_media WHERE title=?',(lesson_title,)).fetchone():
     c.execute('INSERT INTO lesson_media(title,kind,content) VALUES(?,?,?)',(lesson_title,'planner-sample',json.dumps(planner,ensure_ascii=False)))
     added['media']+=1
    storyboard_title='Kịch bản học liệu · An toàn số [MINH HỌA]'
    storyboard={'title':storyboard_title,'note':'Dữ liệu minh họa, không phải video tạo thật','scenes':[{'scene':1,'visual':'Học sinh thảo luận về mật khẩu mạnh','narration':'Không dùng chung mật khẩu.'},{'scene':2,'visual':'Minh họa xác thực hai yếu tố','narration':'Hãy bật xác thực bổ sung khi có thể.'}]}
    if not c.execute('SELECT id FROM lesson_media WHERE title=?',(storyboard_title,)).fetchone():
     c.execute('INSERT INTO lesson_media(title,kind,content) VALUES(?,?,?)',(storyboard_title,'storyboard',json.dumps(storyboard,ensure_ascii=False)))
     added['media']+=1
   return self.sendj(200,{'ok':True,'added':added,'subjects':len(subjects),'notice':'CHỈ LÀ DỮ LIỆU MINH HỌA: Không tạo bài làm, điểm số hay sản phẩm AI giả. Game cần câu hỏi mẫu đa môn đã được tạo trước.'})
  if p=='/api/demo/seed':
   with db() as c:
    if c.execute("SELECT 1 FROM demo_meta WHERE k='seeded'").fetchone():return self.sendj(200,{'ok':True,'already':True})
    classes=[('9A - Minh họa','9'),('8B - Minh họa','8')]
    ids=[]
    for name,grade in classes:
     c.execute('INSERT OR IGNORE INTO classes(name,grade) VALUES(?,?)',(name,grade))
     ids.append(c.execute('SELECT id FROM classes WHERE name=?',(name,)).fetchone()[0])
    sample=[('DEMO901','Nguyễn Minh Anh',ids[0]),('DEMO902','Trần Hải Nam',ids[0]),('DEMO903','Lê Thu Hà',ids[0]),('DEMO904','Phạm Gia Bảo',ids[0]),('DEMO905','Võ Ngọc Linh',ids[0]),('DEMO801','Đỗ Minh Khang',ids[1]),('DEMO802','Hoàng Thảo Vy',ids[1]),('DEMO803','Bùi Đức An',ids[1])]
    for code,name,cid in sample:c.execute('INSERT OR IGNORE INTO students(code,name,class_id) VALUES(?,?,?)',(code,name,cid))
    day=datetime.now().strftime('%Y-%m-%d')
    for i,(code,_,_) in enumerate(sample):
     sid=c.execute('SELECT id FROM students WHERE code=?',(code,)).fetchone()[0]
     c.execute('INSERT OR IGNORE INTO attendance(student_id,day,status) VALUES(?,?,?)',(sid,day,['present','present','late','present','absent','present','excused','present'][i]))
    demo_q=[
     ('Trong Python, lệnh nào kiểm tra điều kiện?','for','if','def','while','B','Câu lệnh if dùng kiểm tra điều kiện.'),
     ('Hàm nào tạo dãy số nguyên trong Python?','range','input','print','len','A','range tạo dãy giá trị.'),
     ('Chỉ số đầu tiên của list trong Python là gì?','1','-1','0','2','C','List đánh chỉ số từ 0.'),
     ('Tìm kiếm nhị phân yêu cầu dữ liệu thế nào?','Ngẫu nhiên','Được sắp xếp','Chỉ số chẵn','Không trùng','B','Dữ liệu phải được sắp xếp.'),
     ('Giá trị của 3 * 4 bằng bao nhiêu?','7','12','34','1','B','Phép nhân cho kết quả 12.'),
     ('Kiểu dữ liệu đúng/sai trong Python là?','int','str','bool','list','C','bool nhận True hoặc False.'),
     ('Vòng lặp nào duyệt lần lượt từng phần tử?','for','if','def','return','A','for dùng để lặp.'),
     ('Độ phức tạp của duyệt tuần tự n phần tử là?','O(1)','O(n)','O(log n)','O(n²)','B','Mỗi phần tử được xét tối đa một lần.')]
    ids=[]
    for text,a,b,copt,dopt,answer,exp in demo_q:
     row=c.execute('SELECT id FROM questions WHERE content=? AND origin=?',(text,'demo_contest')).fetchone()
     if row:ids.append(row[0]);continue
     cur=c.execute('INSERT INTO questions(subject,grade,topic,content,a,b,c,d,answer,difficulty,explanation,origin) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',('Tin học','9','Thuật toán - minh họa',text,a,b,copt,dopt,answer,'Nhận biết',exp,'demo_contest'))
     ids.append(cur.lastrowid)
    title='Đề luyện tập Tin học 9 [MINH HỌA]'
    eidrow=c.execute('SELECT id FROM exams WHERE title=?',(title,)).fetchone()
    if not eidrow:
     eid=c.execute('INSERT INTO exams(title,subject,grade,duration,published) VALUES(?,?,?,?,?)',(title,'Tin học','9',20,1)).lastrowid
     c.executemany('INSERT INTO exam_questions(exam_id,question_id,ord) VALUES(?,?,?)',[(eid,q,i) for i,q in enumerate(ids)])
    if not c.execute('SELECT 1 FROM games WHERE title=?',('Thử thách thuật toán [MINH HỌA]',)).fetchone():
     c.execute('INSERT INTO games(title,topic,kind,question_ids) VALUES(?,?,?,?)',('Thử thách thuật toán [MINH HỌA]','Tin học','quiz',json.dumps(ids[:6])))
    for title,kind in [('Lật thẻ tri thức [MINH HỌA]','flash'),('Đua tốc độ Python [MINH HỌA]','speed'),('Ô cửa bí mật [MINH HỌA]','mystery'),('Leo thang kiến thức [MINH HỌA]','ladder')]:
     c.execute('INSERT INTO games(title,topic,kind,question_ids) VALUES(?,?,?,?)',(title,'Tin học',kind,json.dumps(ids[:6])))
    c.execute("INSERT INTO demo_meta(k,v) VALUES('seeded','1')")
   return self.sendj(200,{'ok':True,'classes':2,'students':8,'label':'DỮ LIỆU MINH HỌA'})
  if p=='/api/questions':
   required=['subject','grade','topic','content','a','b','c','d','answer']
   if any(not str(d.get(k,'')).strip() for k in required):return self.err(400,'Hãy điền đầy đủ câu hỏi và đáp án.')
   if str(d['answer']).upper() not in 'ABCD' or len(str(d['answer']))!=1:return self.err(400,'Đáp án phải là A, B, C hoặc D.')
   with db() as c:
    cur=c.execute('INSERT INTO questions(subject,grade,topic,content,a,b,c,d,answer,difficulty,explanation,origin) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',tuple(str(d[k]).strip()[:2000] for k in required[:-1])+(d['answer'].upper(),str(d.get('difficulty','Nhận biết'))[:40],str(d.get('explanation',''))[:1000],'manual'))
   return self.sendj(201,{'id':cur.lastrowid})
  if p=='/api/exams':
   title=str(d.get('title','')).strip()[:130];ids=d.get('question_ids',[])
   if not title or not isinstance(ids,list) or not 1<=len(ids)<=100:return self.err(400,'Nhập tên đề và chọn 1–100 câu hỏi.')
   try:ids=list(dict.fromkeys(int(x) for x in ids)); duration=int(d.get('duration',45))
   except:return self.err(400,'Dữ liệu câu hỏi hoặc thời lượng sai.')
   if not 1<=duration<=240:return self.err(400,'Thời gian từ 1 đến 240 phút.')
   with db() as c:
    qs=c.execute(f'SELECT id FROM questions WHERE id IN ({",".join("?" for _ in ids)})',ids).fetchall()
    if len(qs)!=len(ids):return self.err(400,'Có câu hỏi không tồn tại.')
    rng=random.SystemRandom();rng.shuffle(ids)
    cur=c.execute('INSERT INTO exams(title,subject,grade,duration,published) VALUES(?,?,?,?,?)',(title,str(d.get('subject','Tin học'))[:80],str(d.get('grade','9'))[:20],duration,1 if d.get('published') else 0))
    c.executemany('INSERT INTO exam_questions(exam_id,question_id,ord) VALUES(?,?,?)',[(cur.lastrowid,q,i) for i,q in enumerate(ids)])
    eid=cur.lastrowid
   return self.sendj(201,{'id':eid})
  if p=='/api/games':
   title=str(d.get('title','')).strip()[:130];topic=str(d.get('topic','')).strip()[:130];kind=str(d.get('kind','quiz'))
   try:ids=list(dict.fromkeys(int(x) for x in d.get('question_ids',[])))
   except:return self.err(400,'Danh sách câu hỏi không hợp lệ.')
   if not title or not topic or kind not in ('quiz','flash','speed','mystery','ladder') or not 1<=len(ids)<=30:return self.err(400,'Cần tên, chủ đề và 1–30 câu hỏi.')
   with db() as c:
    found=c.execute('SELECT COUNT(*) FROM questions WHERE id IN ('+','.join('?' for _ in ids)+')',ids).fetchone()[0]
    if found!=len(ids):return self.err(400,'Câu hỏi không tồn tại.')
    cur=c.execute('INSERT INTO games(title,topic,kind,question_ids) VALUES(?,?,?,?)',(title,topic,kind,json.dumps(ids)))
   return self.sendj(201,{'id':cur.lastrowid})
  if p=='/api/ai/storyboard':
   topic=str(d.get('topic','')).strip()[:200]
   if not topic:return self.err(400,'Hãy nhập chủ đề bài giảng.')
   try:
    result=gemini_text('Bạn là giáo viên. Tạo kịch bản video giáo dục 60-90 giây bằng tiếng Việt, phù hợp học sinh, không dùng thông tin cá nhân. JSON object gồm title (chuỗi), narration (chuỗi), scenes (mảng 4-6 phần tử, mỗi phần tử có visual, voiceover, seconds). Chủ đề: '+topic)
    if not isinstance(result,dict) or not isinstance(result.get('scenes'),list):raise ValueError()
    content=json.dumps(result,ensure_ascii=False)
    with db() as c:cur=c.execute('INSERT INTO lesson_media(title,kind,content) VALUES(?,?,?)',(str(result.get('title',topic))[:130],'storyboard',content))
    return self.sendj(200,{'id':cur.lastrowid,'storyboard':result,'note':'Đây là kịch bản, KHÔNG phải tệp video. Chức năng tạo video thật chưa được kết nối.'})
   except (RuntimeError,ValueError,KeyError,TypeError) as e:return self.err(502,str(e) if isinstance(e,RuntimeError) else 'Gemini trả kịch bản không hợp lệ.')
  if p in ('/api/ai/image','/api/ai/video'):
   media_type='image' if p.endswith('/image') else 'video'
   prompt=str(d.get('prompt','')).strip()[:600]
   if len(prompt)<8:return self.err(400,'Mô tả cần ít nhất 8 ký tự.')
   model=IMAGE_MODEL if media_type=='image' else VIDEO_MODEL
   try:
    description='Tạo một minh họa giáo dục an toàn, không có thông tin cá nhân học sinh: ' if media_type=='image' else 'Tạo video ngắn phục vụ dạy học, không sử dụng hình ảnh hay thông tin nhận dạng trẻ em thật: '
    raw,mime=gemini_interaction(model,description+prompt,media_type,timeout=150)
    allowed={'image':{'image/png':'png','image/jpeg':'jpg','image/webp':'webp'},'video':{'video/mp4':'mp4'}}[media_type]
    ext=allowed.get(mime)
    if not ext or not raw or len(raw)>30_000_000:raise RuntimeError('Tệp trả về không hợp lệ, quá lớn hoặc sai định dạng.')
    filename=uuid.uuid4().hex+'.'+ext;(MEDIA/filename).write_bytes(raw)
    with db() as c:cur=c.execute('INSERT INTO lesson_media(title,kind,content,filename) VALUES(?,?,?,?)',(prompt[:130],media_type,prompt,filename))
    return self.sendj(200,{'id':cur.lastrowid,'url':'/media/'+filename,'kind':media_type,'note':'Nội dung AI tạo thực tế; tệp trên Render Free có thể mất khi khởi động lại.'})
   except (RuntimeError,ValueError,KeyError,TypeError) as e:return self.err(502,str(e) if isinstance(e,RuntimeError) else 'Gemini trả tệp không hợp lệ.')
  if p=='/api/bank16/seed':
   # One subject per request avoids long transactions and proxy 502 timeouts.
   # Retrying a successful request is safe: draft title is the idempotency key.
   from bank16 import FACTS,build_exam
   subject=d.get('subject')
   if not isinstance(subject,str) or subject not in FACTS:
    return self.err(400,'Chọn đúng một môn học cho mỗi lượt tạo đề mẫu.')
   qs,essays=build_exam(subject)
   if len(qs)!=16 or len(essays)!=3:
    return self.err(503,'Bộ câu hỏi minh họa của môn học chưa đủ 16 trắc nghiệm và 3 tự luận.')
   title='Đề ôn tập '+subject+' lớp 8 – 16 TN + 3 TL [MINH HỌA]'
   created={'subjects':0,'questions':0,'draft_exams':0,'essays':0}
   with db() as c:
    c.execute('CREATE TABLE IF NOT EXISTS essay_questions(id INTEGER PRIMARY KEY,subject TEXT NOT NULL,grade TEXT NOT NULL,exam_id INTEGER,content TEXT NOT NULL,origin TEXT NOT NULL)')
    row=c.execute('SELECT id FROM exams WHERE title=?',(title,)).fetchone()
    if row:
     return self.sendj(200,{'subject':subject,'already':True,'created':created,'note':'Đề môn này đã tồn tại; không tạo trùng hoặc sửa dữ liệu cũ.'})
    ids=[]
    for q in qs:
     ids.append(c.execute('INSERT INTO questions(subject,grade,topic,content,a,b,c,d,answer,difficulty,explanation,origin) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',tuple(q[k] for k in ('subject','grade','topic','content','a','b','c','d','answer','difficulty','explanation','origin'))).lastrowid)
    eid=c.execute('INSERT INTO exams(title,subject,grade,duration,published) VALUES(?,?,?,?,0)',(title,subject,'8',45)).lastrowid
    c.executemany('INSERT INTO exam_questions(exam_id,question_id,ord) VALUES(?,?,?)',[(eid,qid,i) for i,qid in enumerate(ids)])
    c.executemany('INSERT INTO essay_questions(subject,grade,exam_id,content,origin) VALUES(?,?,?,?,?)',[(subject,'8',eid,e,'demo-16x19') for e in essays])
    created={'subjects':1,'questions':16,'draft_exams':1,'essays':3}
   return self.sendj(200,{'subject':subject,'already':False,'created':created,'note':'DỮ LIỆU MINH HỌA: đề nháp; tự luận do giáo viên chấm, không tự công bố.'})
  if p=='/api/bank16/essays':
   with db() as c:
    c.execute('CREATE TABLE IF NOT EXISTS essay_questions(id INTEGER PRIMARY KEY,subject TEXT NOT NULL,grade TEXT NOT NULL,exam_id INTEGER,content TEXT NOT NULL,origin TEXT NOT NULL)')
    items=[dict(r) for r in c.execute('SELECT id,subject,grade,exam_id,content FROM essay_questions ORDER BY subject,id LIMIT 200')]
   return self.sendj(200,{'items':items})
  if p=='/api/bank16/ai-essays':
   subject=str(d.get('subject','Toán'))[:70];topic=str(d.get('topic','Kiến thức cơ bản'))[:130];grade=str(d.get('grade','8'))[:12]
   try:
    data=gemini_text('Tạo đúng 3 câu hỏi TỰ LUẬN tiếng Việt cho môn '+subject+', lớp '+grade+', chủ đề '+topic+'. Chỉ trả JSON object {"essays":[{"question":"...","suggested_answer":"..."}]}. Câu hỏi cần phù hợp học sinh; không có dữ liệu cá nhân.')
    if not isinstance(data,dict) or not isinstance(data.get('essays'),list):raise ValueError()
    items=[{'question':str(q['question'])[:1200],'suggested_answer':str(q.get('suggested_answer',''))[:2000]} for q in data['essays'][:3] if isinstance(q,dict) and str(q.get('question','')).strip()]
    if not items:return self.err(502,'Gemini không trả câu tự luận hợp lệ.')
    return self.sendj(200,{'items':items,'note':'Bản nháp AI, cần giáo viên thẩm định; không tự chấm.'})
   except (RuntimeError,ValueError,KeyError,TypeError) as e:return self.err(502,str(e) if isinstance(e,RuntimeError) else 'Gemini trả dữ liệu tự luận không hợp lệ.')
  if p=='/api/ai/generate':
   if not AI_KEY:return self.err(503,'AI chưa cấu hình. Chức năng ngân hàng đề thủ công và chấm trắc nghiệm vẫn hoạt động.')
   subject=str(d.get('subject','Tin học'))[:70];grade=str(d.get('grade','9'))[:20];topic=str(d.get('topic','Thuật toán'))[:130]
   try:n=int(d.get('count',5))
   except:n=5
   if not 1<=n<=10:return self.err(400,'Chỉ tạo 1–10 câu mỗi lần.')
   try:
    data=gemini_text('Tạo '+str(n)+' câu hỏi trắc nghiệm tiếng Việt cho giáo viên, môn '+subject+', lớp '+grade+', chủ đề '+topic+'. Đảm bảo đáp án và giải thích chính xác. Chỉ trả JSON object với questions là danh sách câu hỏi có content,a,b,c,d,answer (A/B/C/D),difficulty,explanation. Không chứa tên học sinh.')
    # Gemini sometimes returns a top-level JSON array instead of {"questions": [...]}.
    # Validate before accessing object methods to avoid aborting the HTTP connection.
    if isinstance(data,list):found=data
    elif isinstance(data,dict):found=data.get('questions',[])
    else:raise ValueError('Cấu trúc phản hồi Gemini không hợp lệ.')
    out=[]
    if not isinstance(found,list):raise ValueError('Danh sách câu hỏi Gemini không hợp lệ.')
    for x in found[:n]:
     if not isinstance(x,dict) or not all(str(x.get(k,'')).strip() for k in ('content','a','b','c','d','answer')) or str(x['answer']).upper() not in ('A','B','C','D'):continue
     out.append({k:str(x.get(k,''))[:2000] for k in ('content','a','b','c','d','answer','difficulty','explanation')})
    if not out:return self.err(502,'Gemini chưa tạo được câu hỏi hợp lệ.')
    return self.sendj(200,{'items':out,'notice':'Câu hỏi Gemini cần được giáo viên duyệt trước khi sử dụng.'})
   except (RuntimeError,ValueError,KeyError,TypeError) as e:return self.err(502,str(e) if isinstance(e,RuntimeError) else 'Gemini trả dữ liệu không hợp lệ.')
  return self.err(404,'Không tìm thấy chức năng.')

if __name__=='__main__':
 if not SECRET:print('CẢNH BÁO: đặt SMARTCLASS_SECRET dài ngẫu nhiên trước khi sử dụng trên internet.')
 if not ADMIN_PASSWORD:print('CẢNH BÁO: chưa đặt SMARTCLASS_ADMIN_PASSWORD, không thể đăng nhập quản trị mới.')
 init();port=int(os.environ.get('PORT','10000'));print(f'SmartClass AI Web listening on {port}',flush=True);ThreadingHTTPServer(('0.0.0.0',port),H).serve_forever()
