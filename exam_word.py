"""Strict and teacher-reviewed DOCX multiple-choice exam importer.

Supports questions starting "Câu 1." or "1.", A/B/C/D options and either
per-question "Đáp án: B" or an answer key section. Plain text only: refuse
graphics, formulas, merged constructs that could silently corrupt exam content.
"""
import io
import re
import zipfile
import secrets
import random
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

Q_RE=re.compile(r'^(?:Câu\s*)?(\d{1,3})\s*[.):-]\s*(.+)$',re.I)
OPT_RE=re.compile(r'^([ABCD])\s*[.):-]\s*(.+)$',re.I)
ANSWER_RE=re.compile(r'^(?:Đáp\s*án|ĐA|Answer)\s*[:：-]\s*([ABCD])\s*$',re.I)
KEY_LINE=re.compile(r'(?<!\w)(\d{1,3})\s*[.):-]\s*([ABCD])(?!\w)',re.I)
UNSUPPORTED=('w:drawing','w:pict','m:oMath','m:oMathPara','w:object','w:altChunk')
MAX_QUESTIONS=150

def parse_word(raw):
 if len(raw)>2_000_000:raise ValueError('File Word tối đa 2 MB.')
 if not zipfile.is_zipfile(io.BytesIO(raw)):raise ValueError('Chỉ chấp nhận file Word .docx hợp lệ.')
 doc=Document(io.BytesIO(raw))
 if any(node.tag.rsplit('}',1)[-1] in ('drawing','pict','oMath','oMathPara','object','altChunk') for node in doc.element.body.iter()):
  raise ValueError('Đề có hình ảnh, công thức Equation hoặc đối tượng nhúng. Chức năng trộn văn bản chưa bảo toàn được các thành phần này; vui lòng dùng đề chỉ có chữ hoặc biên soạn lại để tránh mất nội dung.')
 if doc.tables:
  raise ValueError('Đề có bảng Word. Vui lòng chuyển câu hỏi và đáp án thành các đoạn văn bản trước khi nhập, để tránh trộn sai.')
 lines=[]
 for paragraph in doc.paragraphs:
  for line in paragraph.text.splitlines():
   text=line.strip()
   if text:lines.append(text)
 questions=[];current=None;mode='question';key_text=[]
 for line in lines:
  if re.match(r'^(?:BẢNG\s+)?ĐÁP\s*ÁN(?:\s+TRẮC\s*NGHIỆM)?\s*[:：-]?$',line,re.I) or re.match(r'^ANSWER\s+KEY\s*[:：-]?$',line,re.I):
   mode='key';continue
  if mode=='key':key_text.append(line);continue
  answer=ANSWER_RE.fullmatch(line)
  if answer and current:
   if current['answer'] and current['answer']!=answer.group(1).upper():raise ValueError('Câu có nhiều đáp án không thống nhất.')
   current['answer']=answer.group(1).upper();continue
  option=OPT_RE.match(line)
  if option and current:
   label=option.group(1).upper()
   if label in current['options']:raise ValueError('Câu '+str(current['number'])+' bị trùng phương án '+label+'.')
   current['options'][label]=option.group(2).strip();continue
  match=Q_RE.match(line)
  if match and not (current and len(current['options'])<4 and re.match(r'^\d+\.',line)):
   if current:questions.append(current)
   current={'number':int(match.group(1)),'content':match.group(2).strip(),'options':{},'answer':None}
   if len(questions)>=MAX_QUESTIONS:raise ValueError('Tối đa 150 câu hỏi cho mỗi lần trộn.')
   continue
  if current:
   # Continuation of option/question is safe only when clearly positioned.
   if current['options']:
    last=list(current['options'])[-1]
    current['options'][last]+=' '+line
   else:current['content']+=' '+line
 if current:questions.append(current)
 if not questions:raise ValueError('Không nhận ra câu hỏi. Dùng dạng "Câu 1. Nội dung" và các dòng "A. ...", "B. ...", "C. ...", "D. ...".')
 key={}
 for line in key_text:
  for n,letter in KEY_LINE.findall(line):
   idx=int(n);val=letter.upper()
   if idx in key and key[idx]!=val:raise ValueError('Bảng đáp án bị mâu thuẫn ở câu '+n+'.')
   key[idx]=val
 seen=set()
 for q in questions:
  n=q['number']
  if n in seen:raise ValueError('Số câu bị trùng: '+str(n))
  seen.add(n)
  if set(q['options'])!=set('ABCD'):raise ValueError('Câu '+str(n)+' phải có đủ phương án A, B, C, D.')
  if n in key:
   if q['answer'] and q['answer']!=key[n]:raise ValueError('Đáp án câu '+str(n)+' mâu thuẫn giữa đề và bảng đáp án.')
   q['answer']=key[n]
  if not q['answer']:raise ValueError('Thiếu đáp án câu '+str(n)+'. Thêm "Đáp án: A" hoặc bảng ĐÁP ÁN ở cuối file.')
  if any(not v.strip() for v in q['options'].values()) or not q['content'].strip():raise ValueError('Câu '+str(n)+' có nội dung trống.')
 return questions

def validate_questions(items):
 if not isinstance(items,list) or not 1<=len(items)<=MAX_QUESTIONS:raise ValueError('Cần từ 1 đến 150 câu hợp lệ.')
 clean=[]
 for n,q in enumerate(items,1):
  if not isinstance(q,dict):raise ValueError('Dữ liệu câu hỏi không hợp lệ.')
  text=str(q.get('content','')).strip()
  options=q.get('options')
  answer=str(q.get('answer','')).upper().strip()
  if len(text)<3 or len(text)>4000 or not isinstance(options,dict) or set(options)!=set('ABCD') or answer not in 'ABCD':
   raise ValueError('Câu '+str(n)+' thiếu nội dung, phương án hoặc đáp án.')
  opts={key:str(options[key]).strip() for key in 'ABCD'}
  if any(not x or len(x)>2000 for x in opts.values()):raise ValueError('Phương án không hợp lệ ở câu '+str(n)+'.')
  clean.append({'number':n,'content':text,'options':opts,'answer':answer})
 return clean

def make_variants(items,amount):
 if amount not in (2,4,8):raise ValueError('Chỉ hỗ trợ 2, 4 hoặc 8 mã đề.')
 questions=validate_questions(items)
 variants=[];seen=set()
 for k in range(amount):
  for attempt in range(120):
   rng=random.Random(secrets.randbits(128))
   sequence=list(range(len(questions)));rng.shuffle(sequence)
   entries=[];signature=[]
   for idx in sequence:
    q=questions[idx];letters=list('ABCD');rng.shuffle(letters)
    choices={new:q['options'][old] for new,old in zip('ABCD',letters)}
    correct='ABCD'[letters.index(q['answer'])]
    entries.append({'source_number':q['number'],'content':q['content'],'options':choices,'answer':correct})
    signature.append((idx,tuple(letters)))
   signature=tuple(signature)
   if signature not in seen:
    seen.add(signature);variants.append(entries);break
  else:raise ValueError('Không thể tạo đủ mã đề khác nhau; hãy bổ sung câu hoặc phương án.')
 return variants

def build_doc(title,code,questions):
 doc=Document();section=doc.sections[0]
 section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(2.3);section.right_margin=Cm(2)
 normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
 h=doc.add_heading(title,level=1);h.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('MÃ ĐỀ: '+code).bold=True
 doc.add_paragraph('Họ và tên: .............................................    Lớp: .......................')
 for i,q in enumerate(questions,1):
  p=doc.add_paragraph(style='Normal');p.paragraph_format.keep_with_next=True
  p.add_run(f'Câu {i}. ').bold=True;p.add_run(q['content'])
  for letter in 'ABCD':
   p=doc.add_paragraph(letter+'. '+q['options'][letter]);p.paragraph_format.left_indent=Cm(0.7)
 output=io.BytesIO();doc.save(output);return output.getvalue()

def export_package(title,questions,count):
 if len(title.strip())<3 or len(title)>150:raise ValueError('Tên đề cần từ 3 đến 150 ký tự.')
 variants=make_variants(questions,count)
 package=io.BytesIO()
 with zipfile.ZipFile(package,'w',zipfile.ZIP_DEFLATED) as archive:
  answer_doc=Document();answer_doc.add_heading('BẢNG ĐÁP ÁN - '+title,0)
  answer_doc.add_paragraph('Đáp án được cập nhật theo thứ tự câu và phương án sau khi trộn.')
  for i,variant in enumerate(variants,1):
   code=f'{i:03d}'
   archive.writestr(f'Ma_de_{code}.docx',build_doc(title,code,variant))
   answer_doc.add_heading('Mã đề '+code,level=2)
   answer_doc.add_paragraph('  '.join(f'{j}.{q["answer"]}' for j,q in enumerate(variant,1)))
  buffer=io.BytesIO();answer_doc.save(buffer)
  archive.writestr('Bang_dap_an.docx',buffer.getvalue())
 return package.getvalue()
